import json
import os
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
import re

from docx import Document
from dotenv import load_dotenv
from flask import Flask, flash, redirect, render_template, request, send_file, url_for
import httpx
from openai import OpenAI
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)
from xml.sax.saxutils import escape
from werkzeug.utils import secure_filename


BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
OUTPUT_DIR = BASE_DIR / "output"
ALLOWED_EXTENSIONS = {".docx"}

load_dotenv()

client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY"),
)

app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("FLASK_SECRET_KEY", "dev-secret-change-me")
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024

UPLOAD_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)


RESULT_SECTIONS = [
    ("evaluasi_judul", "Evaluasi Judul"),
    ("rekomendasi_judul_baru", "Rekomendasi Judul Baru"),
    ("evaluasi_abstrak", "Evaluasi Abstrak"),
    ("abstrak_versi_perbaikan", "Abstrak Versi Perbaikan"),
    ("evaluasi_pendahuluan", "Evaluasi Pendahuluan"),
    ("gap_penelitian", "Gap Penelitian"),
    ("cek_kutipan_dan_daftar_pustaka", "Cek Kutipan dan Daftar Pustaka"),
    ("rumusan_masalah", "Rumusan Masalah"),
    ("pendahuluan_versi_revisi", "Pendahuluan Versi Revisi"),
]

TITLE_SECTIONS = [
    ("judul_terdeteksi", "Judul Terdeteksi"),
    ("evaluasi_judul", "Evaluasi Judul"),
    ("masalah_judul", "Masalah Judul"),
    ("kekuatan_judul", "Kekuatan Judul"),
    ("rekomendasi_judul_baru", "Rekomendasi Judul Baru"),
    ("skor_judul", "Skor Judul"),
]

ABSTRACT_SECTIONS = [
    ("abstrak_terdeteksi", "Abstrak Terdeteksi"),
    ("evaluasi_abstrak", "Evaluasi Abstrak"),
    ("kekuatan_abstrak", "Kekuatan Abstrak"),
    ("kelemahan_abstrak", "Kelemahan Abstrak"),
    ("kelengkapan_unsur_abstrak", "Kelengkapan Unsur Abstrak"),
    ("saran_perbaikan", "Saran Perbaikan"),
    ("abstrak_versi_perbaikan", "Abstrak Versi Perbaikan"),
    ("skor_abstrak", "Skor Abstrak"),
]

INTRODUCTION_SECTIONS = [
    ("evaluasi_pendahuluan", "Evaluasi Pendahuluan"),
    ("kekuatan_pendahuluan", "Kekuatan Pendahuluan"),
    ("kelemahan_pendahuluan", "Kelemahan Pendahuluan"),
    ("gap_penelitian", "Gap Penelitian"),
    ("rumusan_masalah_dan_tujuan", "Rumusan Masalah dan Tujuan"),
    ("evaluasi_kutipan_pendahuluan", "Evaluasi Kutipan Pendahuluan"),
    ("saran_perbaikan", "Saran Perbaikan"),
    ("pendahuluan_versi_revisi", "Pendahuluan Versi Revisi"),
    ("skor_pendahuluan", "Skor Pendahuluan"),
]

METHODOLOGY_SECTIONS = [
    ("evaluasi_umum_metodologi", "Evaluasi Umum Metodologi"),
    ("kelemahan_metodologi", "Kelemahan Metodologi"),
    (
        "kesesuaian_metode_dengan_tujuan_rumusan_masalah",
        "Kesesuaian Metode dengan Tujuan/Rumusan Masalah",
    ),
    ("komponen_metodologi_yang_belum_ada", "Komponen Metodologi yang Belum Ada"),
    ("saran_perbaikan", "Saran Perbaikan"),
    ("metodologi_versi_revisi", "Metodologi Versi Revisi"),
    ("catatan_yang_perlu_dilengkapi_penulis", "Catatan yang Perlu Dilengkapi Penulis"),
]

RESULTS_DISCUSSION_SECTIONS = [
    ("evaluasi_hasil_penelitian", "Evaluasi Hasil Penelitian"),
    ("evaluasi_pembahasan", "Evaluasi Pembahasan"),
    ("kesesuaian_dengan_rumusan_masalah", "Kesesuaian dengan Rumusan Masalah"),
    ("kesesuaian_dengan_metodologi", "Kesesuaian dengan Metodologi"),
    ("kesesuaian_dengan_tujuan_penelitian", "Kesesuaian dengan Tujuan Penelitian"),
    ("inkonsistensi_yang_ditemukan", "Inkonsistensi yang Ditemukan"),
    ("saran_perbaikan", "Saran Perbaikan"),
    ("hasil_versi_revisi", "Hasil Versi Revisi"),
    ("pembahasan_versi_revisi", "Pembahasan Versi Revisi"),
    ("catatan_untuk_penulis", "Catatan untuk Penulis"),
]

CONSISTENCY_SECTIONS = [
    ("ringkasan_kualitas_konsistensi_artikel", "Ringkasan Kualitas Konsistensi Artikel"),
    ("skor_konsistensi_artikel", "Skor Konsistensi Artikel 1-100"),
    ("peta_keselarasan_artikel", "Peta Keselarasan Artikel"),
    ("daftar_inkonsistensi_yang_ditemukan", "Daftar Inkonsistensi yang Ditemukan"),
    ("dampak_akademik_dari_inkonsistensi", "Dampak Akademik dari Inkonsistensi"),
    ("prioritas_perbaikan", "Prioritas Perbaikan"),
    ("rekomendasi_revisi_menyeluruh", "Rekomendasi Revisi Menyeluruh"),
    ("checklist_kesiapan_artikel_untuk_submit_jurnal", "Checklist Kesiapan Artikel untuk Submit Jurnal"),
]

REVIEWER_SIMULATOR_SECTIONS = [
    ("kelebihan_artikel", "Kelebihan Artikel"),
    ("kelemahan_artikel", "Kelemahan Artikel"),
    ("novelty_penelitian", "Novelty Penelitian"),
    ("signifikansi_penelitian", "Signifikansi Penelitian"),
    ("kekuatan_metodologi", "Kekuatan Metodologi"),
    ("kelemahan_metodologi", "Kelemahan Metodologi"),
    ("kualitas_pembahasan", "Kualitas Pembahasan"),
    ("kualitas_referensi", "Kualitas Referensi"),
    ("kualitas_bahasa_akademik", "Kualitas Bahasa Akademik"),
    ("potensi_publikasi", "Potensi Publikasi"),
    ("keputusan_reviewer", "Keputusan Reviewer"),
    ("alasan_keputusan", "Alasan Keputusan"),
]

REFERENCE_CITATION_SECTIONS = [
    ("ringkasan_kualitas_referensi", "Ringkasan Kualitas Referensi"),
    ("statistik_referensi", "Statistik Referensi"),
    ("statistik_usia_referensi", "Statistik Usia Referensi"),
    ("sitasi_bermasalah", "Sitasi Bermasalah"),
    ("referensi_bermasalah", "Referensi Bermasalah"),
    ("referensi_tidak_pernah_disitasi", "Referensi Tidak Pernah Disitasi"),
    ("sitasi_tanpa_referensi", "Sitasi Tanpa Referensi"),
    ("referensi_duplikat", "Referensi Duplikat"),
    ("evaluasi_kualitas_referensi", "Evaluasi Kualitas Referensi"),
    ("evaluasi_kepatuhan_apa_7", "Evaluasi Kepatuhan APA 7"),
    ("daftar_perbaikan_yang_direkomendasikan", "Daftar Perbaikan yang Direkomendasikan"),
    ("skor_kualitas_referensi", "Skor Kualitas Referensi (0-100)"),
    ("tingkat_kesiapan_referensi_untuk_publikasi", "Tingkat Kesiapan Referensi untuk Publikasi"),
]

NOVELTY_CONTRIBUTION_SECTIONS = [
    ("ringkasan_novelty_penelitian", "Ringkasan Novelty Penelitian"),
    ("ringkasan_kontribusi_ilmiah", "Ringkasan Kontribusi Ilmiah"),
    ("evaluasi_gap_penelitian", "Evaluasi Gap Penelitian"),
    ("novelty_yang_teridentifikasi", "Novelty yang Teridentifikasi"),
    ("kontribusi_teoretis", "Kontribusi Teoretis"),
    ("kontribusi_praktis", "Kontribusi Praktis"),
    ("kontribusi_metodologis", "Kontribusi Metodologis"),
    ("risiko_akademik", "Risiko Akademik"),
    ("saran_penguatan_novelty", "Saran Penguatan Novelty"),
    ("saran_penguatan_kontribusi", "Saran Penguatan Kontribusi"),
    ("skor_novelty", "Skor Novelty (0-100)"),
    ("skor_kontribusi_ilmiah", "Skor Kontribusi Ilmiah (0-100)"),
    ("tingkat_kebaruan_penelitian", "Tingkat Kebaruan Penelitian (Level 1-7)"),
    ("kelayakan_kontribusi_untuk_publikasi", "Kelayakan Kontribusi untuk Publikasi"),
]

SUBMISSION_READINESS_SECTIONS = [
    ("ringkasan_evaluasi_artikel", "Ringkasan Evaluasi Artikel"),
    ("checklist_kesiapan_submit", "Checklist Kesiapan Submit"),
    ("skor_kesiapan_submit", "Skor Kesiapan Submit"),
    ("kategori_kelayakan", "Kategori Kelayakan"),
    ("kelebihan_artikel", "Kelebihan Artikel"),
    ("kelemahan_artikel", "Kelemahan Artikel"),
    ("risiko_penolakan", "Risiko Penolakan"),
    ("prioritas_revisi", "Prioritas Revisi"),
    ("rekomendasi_akhir_editor", "Rekomendasi Akhir Editor"),
]

PDF_REVIEW_SECTIONS = [
    ("judul", "Judul", [("judul_awal", "Judul Awal"), ("judul_rekomendasi", "Judul Rekomendasi")]),
    ("abstrak", "Abstrak Revisi", [("abstrak_revisi", "Abstrak Revisi")]),
    ("pendahuluan", "Pendahuluan Revisi", [("pendahuluan_revisi", "Pendahuluan Revisi")]),
    ("metodologi", "Metodologi Revisi", [("metodologi_revisi", "Metodologi Revisi")]),
    (
        "hasil_pembahasan",
        "Hasil dan Pembahasan Revisi",
        [("hasil_pembahasan_revisi", "Hasil dan Pembahasan Revisi")],
    ),
    ("kesimpulan", "Kesimpulan Revisi", [("kesimpulan_revisi", "Kesimpulan Revisi")]),
    (
        "sitasi_apa",
        "Sitasi dan APA 7",
        [
            ("sitasi_tanpa_referensi", "Sitasi Tanpa Referensi"),
            ("referensi_tidak_disitasi", "Referensi Tidak Disitasi"),
            ("referensi_bermasalah", "Referensi Bermasalah"),
            ("catatan_akhir_sitasi", "Catatan Akhir Sitasi"),
        ],
    ),
    (
        "referensi_tambahan",
        "Referensi Tambahan",
        [
            ("referensi_asli", "A. Referensi Asli Penulis"),
            ("rekomendasi", "B. Referensi Tambahan yang Direkomendasikan"),
            ("referensi_dipilih", "C. Referensi Tambahan Terpilih dalam Format APA 7"),
        ],
    ),
]

WORKFLOW_STEPS = [
    {
        "number": 1,
        "key": "judul",
        "endpoint": "analyze_title",
        "label": "Judul",
    },
    {
        "number": 2,
        "key": "abstrak",
        "endpoint": "analyze_abstract",
        "label": "Abstrak",
    },
    {
        "number": 3,
        "key": "pendahuluan",
        "endpoint": "analyze_introduction",
        "label": "Pendahuluan",
    },
    {
        "number": 4,
        "key": "metodologi",
        "endpoint": "analyze_methodology",
        "label": "Metodologi",
    },
    {
        "number": 5,
        "key": "hasil_pembahasan",
        "endpoint": "analyze_results_discussion",
        "label": "Hasil dan Pembahasan",
    },
    {
        "number": 6,
        "key": "kesimpulan",
        "endpoint": "analyze_conclusion",
        "label": "Kesimpulan",
    },
    {
        "number": 7,
        "key": "sitasi_apa",
        "endpoint": "analyze_reference_citation",
        "label": "Sitasi dan APA 7",
    },
    {
        "number": 8,
        "key": "referensi_tambahan",
        "endpoint": "analyze_reference_recommendations",
        "label": "Rekomendasi Referensi Tambahan",
    },
]

METHODOLOGY_HEADINGS = {
    "metode",
    "metodologi",
    "metode penelitian",
    "desain penelitian",
    "prosedur penelitian",
    "metode dan prosedur penelitian",
    "method",
    "methods",
    "methodology",
    "research method",
    "research methods",
    "materials and methods",
    "methods and materials",
    "research design",
    "research procedure",
    "research procedures",
}

RESULTS_DISCUSSION_HEADINGS = {
    "hasil",
    "hasil penelitian",
    "results",
    "result",
    "pembahasan",
    "discussion",
    "hasil dan pembahasan",
    "results and discussion",
    "temuan penelitian",
    "findings",
    "research findings",
}

NEXT_SECTION_HEADINGS = {
    "hasil",
    "hasil penelitian",
    "hasil dan pembahasan",
    "pembahasan",
    "results",
    "result",
    "results and discussion",
    "discussion",
    "temuan",
    "temuan penelitian",
    "findings",
    "research findings",
    "kesimpulan",
    "simpulan",
    "conclusion",
    "conclusions",
    "daftar pustaka",
    "referensi",
    "references",
    "bibliography",
}

RESULTS_DISCUSSION_STOP_HEADINGS = {
    "kesimpulan",
    "simpulan",
    "conclusion",
    "conclusions",
    "penutup",
    "closing",
    "saran",
    "recommendation",
    "recommendations",
    "limitations",
    "keterbatasan",
    "daftar pustaka",
    "referensi",
    "references",
    "bibliography",
    "lampiran",
    "appendix",
}

CONCLUSION_HEADINGS = {
    "kesimpulan",
    "simpulan",
    "conclusion",
    "conclusions",
}

CONCLUSION_STOP_HEADINGS = {
    "daftar pustaka",
    "referensi",
    "references",
    "bibliography",
    "lampiran",
    "appendix",
}

REFERENCE_HEADINGS = {
    "daftar pustaka",
    "referensi",
    "references",
    "bibliography",
}

ABSTRACT_HEADINGS = {
    "abstrak",
    "abstract",
}

ABSTRACT_STOP_HEADINGS = {
    "keywords",
    "kata kunci",
    "introduction",
    "pendahuluan",
    "latar belakang",
    "background",
}

INTRODUCTION_HEADINGS = {
    "introduction",
    "pendahuluan",
    "latar belakang",
    "background",
}

INTRODUCTION_STOP_HEADINGS = {
    "research method",
    "research methods",
    "method",
    "methods",
    "methodology",
    "metode",
    "metodologi",
    "metode penelitian",
    "desain penelitian",
    "prosedur penelitian",
    "metode dan prosedur penelitian",
    "materials and methods",
    "methods and materials",
    "research design",
    "research procedure",
    "research procedures",
    # Fallback jika artikel melompati heading metodologi.
    "hasil",
    "hasil penelitian",
    "hasil dan pembahasan",
    "pembahasan",
    "temuan penelitian",
    "kesimpulan",
    "simpulan",
    "daftar pustaka",
    "referensi",
    "references",
    "bibliography",
    "results",
    "discussion",
    "results and discussion",
    "findings",
    "research findings",
    "conclusion",
    "conclusions",
}


def allowed_file(filename):
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def get_openai_warning():
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key or api_key == "isi_api_key":
        return "Peringatan: isi OPENAI_API_KEY yang valid di file .env."
    return None


def extract_docx_text(path):
    document = Document(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    table_text = []

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_text.append(" | ".join(cells))

    text_blocks = [item for item in paragraphs + table_text if item]
    return "\n".join(text_blocks)


def normalize_heading(text):
    normalized = re.sub(r"^\s*(?:bab\s+)?[ivxlcdm\d]+[\.\)]?\s+", "", text, flags=re.IGNORECASE)
    normalized = re.sub(r"[^a-zA-Z\s]", " ", normalized)
    normalized = re.sub(r"\s+", " ", normalized).strip().lower()
    return normalized


def heading_matches(text, keywords):
    normalized_heading = normalize_heading(text)
    return any(keyword in normalized_heading for keyword in keywords)


def detect_title(article_text):
    lines = [line.strip() for line in article_text.splitlines() if line.strip()]
    title_candidates = []
    stopped_at_abstract = False

    stop_heading_found = None
    print("SECTION START:", lines[0] if lines else None)

    for line in lines:
        if heading_matches(line, ABSTRACT_HEADINGS):
            stopped_at_abstract = True
            stop_heading_found = line
            break

        title_candidates.append(line)

    print("SECTION STOP:", stop_heading_found)

    if not title_candidates:
        return {
            "detected": False,
            "title": "",
            "note": "Judul tidak terdeteksi sebelum bagian Abstrak/Abstract.",
        }

    title = title_candidates[0]
    if len(title) < 12 and len(title_candidates) > 1:
        title = " ".join(title_candidates[:2])

    note = "Kandidat judul diambil dari baris awal artikel sebelum bagian Abstrak/Abstract."
    if not stopped_at_abstract:
        note = (
            "Bagian Abstrak/Abstract tidak ditemukan. Kandidat judul tetap diambil "
            "dari baris awal artikel."
        )

    return {
        "detected": True,
        "title": title,
        "note": note,
    }


def split_abstract_keywords(line):
    pattern = r"\b(?:keywords|kata\s+kunci)\b\s*:?"
    parts = re.split(pattern, line, maxsplit=1, flags=re.IGNORECASE)
    return parts[0].strip()


def detect_abstract_section(article_text):
    lines = [line.strip() for line in article_text.splitlines() if line.strip()]
    start_index = None
    detected_heading = None
    inline_abstract_text = ""

    for index, line in enumerate(lines):
        if heading_matches(line, ABSTRACT_HEADINGS):
            start_index = index
            detected_heading = line
            if ":" in line:
                inline_abstract_text = line.split(":", 1)[1].strip()
            break

    if start_index is None:
        print("SECTION START:", None)
        print("SECTION STOP:", None)
        return {
            "detected": False,
            "heading": None,
            "text": "",
            "note": "Bagian abstrak tidak terdeteksi melalui heading Abstract atau Abstrak.",
        }

    print("SECTION START:", detected_heading)
    abstract_lines = []
    if inline_abstract_text:
        cleaned_inline = split_abstract_keywords(inline_abstract_text)
        if cleaned_inline:
            abstract_lines.append(cleaned_inline)

    stop_heading_found = None
    for index in range(start_index + 1, len(lines)):
        line = lines[index]
        if heading_matches(line, ABSTRACT_STOP_HEADINGS):
            stop_heading_found = line
            break

        cleaned_line = split_abstract_keywords(line)
        if cleaned_line != line:
            if cleaned_line:
                abstract_lines.append(cleaned_line)
            break

        abstract_lines.append(line)

    print("SECTION STOP:", stop_heading_found)

    abstract_text = "\n".join(abstract_lines).strip()
    note = "Teks abstrak diambil setelah heading Abstrak/Abstract sampai sebelum kata kunci atau pendahuluan."
    if not abstract_text:
        note = "Heading Abstrak/Abstract ditemukan, tetapi teks abstrak setelah heading tidak terdeteksi."

    return {
        "detected": True,
        "heading": detected_heading,
        "text": abstract_text,
        "note": note,
    }


def detect_introduction_section(article_text):
    lines = [line.strip() for line in article_text.splitlines() if line.strip()]
    start_index = None
    detected_heading = None

    for index, line in enumerate(lines):
        if heading_matches(line, INTRODUCTION_HEADINGS):
            start_index = index
            detected_heading = line
            break

    if start_index is None:
        print("SECTION START:", None)
        print("SECTION STOP:", None)
        return {
            "detected": False,
            "heading": None,
            "text": "",
            "note": (
                "Bagian pendahuluan tidak terdeteksi secara eksplisit melalui heading "
                "Introduction atau Pendahuluan."
            ),
        }

    print("SECTION START:", detected_heading)
    end_index = len(lines)
    stop_heading_found = None
    for index in range(start_index + 1, len(lines)):
        if heading_matches(lines[index], INTRODUCTION_STOP_HEADINGS):
            end_index = index
            stop_heading_found = lines[index]
            break

    print("SECTION STOP:", stop_heading_found)
    section_text = "\n".join(lines[start_index:end_index]).strip()
    note = "Teks pendahuluan diambil dari heading pendahuluan sampai sebelum section berikutnya."
    if end_index == len(lines):
        note = (
            "Heading section berikutnya tidak ditemukan setelah pendahuluan. Teks diambil dari "
            "heading pendahuluan sampai akhir artikel."
        )

    return {
        "detected": True,
        "heading": detected_heading,
        "text": section_text,
        "note": note,
    }


def detect_methodology_section(article_text):
    lines = [line.strip() for line in article_text.splitlines() if line.strip()]
    start_index = None
    detected_heading = None

    for index, line in enumerate(lines):
        if heading_matches(line, METHODOLOGY_HEADINGS):
            start_index = index
            detected_heading = line
            break

    if start_index is None:
        print("SECTION START:", None)
        print("SECTION STOP:", None)
        return {
            "detected": False,
            "heading": None,
            "text": (
                "Bagian metodologi tidak terdeteksi secara eksplisit."
            ),
        }

    print("SECTION START:", detected_heading)
    end_index = len(lines)
    stop_heading_found = None
    for index in range(start_index + 1, len(lines)):
        if heading_matches(lines[index], NEXT_SECTION_HEADINGS):
            end_index = index
            stop_heading_found = lines[index]
            break

    print("SECTION STOP:", stop_heading_found)
    section_text = "\n".join(lines[start_index:end_index]).strip()
    return {
        "detected": True,
        "heading": detected_heading,
        "text": section_text,
    }


def detect_results_discussion_section(article_text):
    lines = [line.strip() for line in article_text.splitlines() if line.strip()]
    start_index = None
    detected_heading = None

    for index, line in enumerate(lines):
        if heading_matches(line, RESULTS_DISCUSSION_HEADINGS):
            start_index = index
            detected_heading = line
            break

    if start_index is None:
        print("SECTION START:", None)
        print("SECTION STOP:", None)
        return {
            "detected": False,
            "heading": None,
            "text": (
                "Bagian hasil dan pembahasan tidak terdeteksi secara eksplisit."
            ),
        }

    print("SECTION START:", detected_heading)
    end_index = len(lines)
    stop_heading_found = None
    for index in range(start_index + 1, len(lines)):
        if heading_matches(lines[index], RESULTS_DISCUSSION_STOP_HEADINGS):
            end_index = index
            stop_heading_found = lines[index]
            break

    print("SECTION STOP:", stop_heading_found)
    section_text = "\n".join(lines[start_index:end_index]).strip()
    return {
        "detected": True,
        "heading": detected_heading,
        "text": section_text,
    }


def detect_section(article_text, start_keywords, stop_keywords, section_name):
    lines = [line.strip() for line in article_text.splitlines() if line.strip()]
    start_index = None
    detected_heading = None

    for index, line in enumerate(lines):
        if heading_matches(line, start_keywords):
            start_index = index
            detected_heading = line
            break

    if start_index is None:
        print("SECTION START:", None)
        print("SECTION STOP:", None)
        return {
            "detected": False,
            "heading": None,
            "text": "",
            "note": f"Bagian {section_name} tidak terdeteksi secara eksplisit.",
        }

    stop_heading_found = None
    end_index = len(lines)
    for index in range(start_index + 1, len(lines)):
        if heading_matches(lines[index], stop_keywords):
            end_index = index
            stop_heading_found = lines[index]
            break

    print("SECTION START:", detected_heading)
    print("SECTION STOP:", stop_heading_found)
    return {
        "detected": True,
        "heading": detected_heading,
        "text": "\n".join(lines[start_index:end_index]).strip(),
        "note": f"Bagian {section_name} diambil sampai sebelum section berikutnya.",
    }


def detect_conclusion_section(article_text):
    return detect_section(
        article_text,
        CONCLUSION_HEADINGS,
        CONCLUSION_STOP_HEADINGS,
        "kesimpulan",
    )


def detect_reference_section(article_text):
    return detect_section(
        article_text,
        REFERENCE_HEADINGS,
        {"lampiran", "appendix"},
        "daftar pustaka",
    )


def build_short_article_context(article_text):
    sections = [
        detect_title(article_text).get("title", ""),
        detect_abstract_section(article_text).get("text", ""),
        detect_introduction_section(article_text).get("text", ""),
    ]
    context = "\n\n".join(item for item in sections if item)
    return context[:5000]


def build_title_prompt(title_text):
    return f"""
Anda adalah reviewer akademik yang khusus menilai judul artikel ilmiah.

Tugas Anda:
1. Evaluasi kejelasan, ketepatan, kebaruan, dan kekuatan akademik judul.
2. Identifikasi masalah judul seperti terlalu umum, terlalu panjang, ambigu, tidak mencerminkan variabel, konteks, metode, atau fokus artikel.
3. Identifikasi kekuatan judul jika ada.
4. Berikan alternatif judul baru yang lebih akademik.
5. Berikan skor judul dari 0 sampai 100.

Aturan wajib:
- Jangan mengarang data penelitian, lokasi, metode, variabel, subjek, atau konteks yang tidak muncul pada judul.
- Jika informasi pada judul terbatas, tulis keterbatasannya secara eksplisit.
- Gunakan bahasa Indonesia akademik, jelas, formal, dan konstruktif.
- Kembalikan jawaban hanya dalam JSON valid.

Format JSON wajib:
{{
  "judul_terdeteksi": "judul artikel",
  "evaluasi_judul": "teks evaluasi",
  "masalah_judul": ["masalah 1", "masalah 2"],
  "kekuatan_judul": ["kekuatan 1", "kekuatan 2"],
  "rekomendasi_judul_baru": ["judul 1", "judul 2", "judul 3"],
  "skor_judul": 0
}}

Judul artikel:
\"\"\"{title_text}\"\"\"
""".strip()


def build_abstract_prompt(abstract_info, article_text):
    abstract_text = abstract_info.get("text") or ""
    detection_note = abstract_info.get("note") or ""

    return f"""
Anda adalah reviewer akademik yang khusus menilai abstrak artikel ilmiah.

Tugas Anda:
1. Evaluasi kualitas abstrak secara akademik.
2. Cek kelengkapan unsur abstrak: tujuan, metode, subjek/partisipan, teknik pengumpulan data, teknik analisis data, hasil, dan kontribusi/implikasi.
3. Cek apakah abstrak terlalu panjang, terlalu umum, atau terlalu banyak klaim.
4. Cek apakah klaim hasil dalam abstrak didukung informasi dalam artikel.
5. Berikan saran perbaikan yang konkret.
6. Buat versi abstrak perbaikan tanpa mengarang data baru.

Aturan wajib:
- Jangan mengarang jumlah sampel, lokasi, angka hasil, skor, instrumen, atau temuan baru.
- Jika informasi tidak ada di artikel, tulis bahwa informasi tersebut perlu dilengkapi penulis.
- Jangan menambahkan sitasi baru.
- Gunakan bahasa Indonesia akademik.
- Kembalikan hanya JSON valid.

Format JSON wajib:
{{
  "abstrak_terdeteksi": "teks abstrak",
  "evaluasi_abstrak": "teks evaluasi",
  "kekuatan_abstrak": ["kekuatan 1", "kekuatan 2"],
  "kelemahan_abstrak": ["kelemahan 1", "kelemahan 2"],
  "kelengkapan_unsur_abstrak": {{
    "tujuan": "ada/tidak ada/perlu diperjelas",
    "metode": "ada/tidak ada/perlu diperjelas",
    "partisipan": "ada/tidak ada/perlu diperjelas",
    "pengumpulan_data": "ada/tidak ada/perlu diperjelas",
    "analisis_data": "ada/tidak ada/perlu diperjelas",
    "hasil": "ada/tidak ada/perlu diperjelas",
    "kontribusi": "ada/tidak ada/perlu diperjelas"
  }},
  "saran_perbaikan": ["saran 1", "saran 2"],
  "abstrak_versi_perbaikan": "versi abstrak perbaikan",
  "skor_abstrak": 0
}}

Catatan deteksi:
\"\"\"{detection_note}\"\"\"

Teks abstrak terdeteksi:
\"\"\"{abstract_text}\"\"\"
""".strip()


def build_introduction_prompt(introduction_info, article_text):
    introduction_text = introduction_info.get("text") or ""
    detection_note = introduction_info.get("note") or ""

    return f"""
Anda adalah reviewer akademik yang khusus menilai bagian pendahuluan artikel ilmiah.

Tugas Anda:
1. Evaluasi kualitas pendahuluan secara akademik.
2. Nilai kekuatan latar belakang penelitian.
3. Nilai kejelasan gap penelitian.
4. Nilai rumusan masalah dan/atau tujuan penelitian.
5. Cek apakah kutipan dalam pendahuluan relevan dengan argumen yang dibangun.
6. Berikan saran perbaikan yang konkret dan akademik.
7. Buat versi revisi pendahuluan tanpa mengarang data, temuan, lokasi, metode, teori, atau sumber baru.

Aturan wajib:
- Jangan mengarang data penelitian, angka, lokasi, metode, temuan, teori, atau konteks yang tidak tersedia.
- Jangan membuat kutipan palsu atau menambahkan sumber baru.
- Jika informasi tidak cukup, tulis keterbatasannya secara eksplisit.
- Revisi pendahuluan hanya boleh memakai informasi yang ada pada artikel.
- Gunakan bahasa Indonesia akademik, jelas, formal, dan konstruktif.
- Kembalikan jawaban hanya dalam JSON valid.

Format JSON wajib:
{{
  "evaluasi_pendahuluan": "teks evaluasi",
  "kekuatan_pendahuluan": ["kekuatan 1", "kekuatan 2"],
  "kelemahan_pendahuluan": ["kelemahan 1", "kelemahan 2"],
  "gap_penelitian": "evaluasi gap penelitian",
  "rumusan_masalah_dan_tujuan": "evaluasi rumusan masalah dan tujuan",
  "evaluasi_kutipan_pendahuluan": ["evaluasi kutipan 1", "evaluasi kutipan 2"],
  "saran_perbaikan": ["saran 1", "saran 2"],
  "pendahuluan_versi_revisi": "versi revisi pendahuluan",
  "skor_pendahuluan": 0
}}

Catatan deteksi:
\"\"\"{detection_note}\"\"\"

Teks pendahuluan terdeteksi:
\"\"\"{introduction_text}\"\"\"
""".strip()


def build_prompt(article_text):
    return f"""
Anda adalah reviewer artikel ilmiah berbahasa Indonesia yang fokus memperbaiki PENDAHULUAN agar matang secara akademik.

Tugas Anda:
1. Nilai kesesuaian judul dengan isi artikel.
2. Deteksi apakah judul ambigu, terlalu umum, kurang akademik, atau belum mencerminkan variabel/konteks.
3. Berikan alternatif judul baru yang akademik.
4. Evaluasi abstrak dan buat versi perbaikannya jika informasi memadai.
5. Evaluasi pendahuluan secara mendalam.
6. Periksa kekuatan latar belakang, kejelasan gap penelitian, dan keberadaan rumusan masalah.
7. Identifikasi kutipan di pendahuluan, cocokkan dengan daftar pustaka yang tersedia di teks artikel.
8. Tandai kutipan yang tidak ditemukan di daftar pustaka dengan frasa "perlu verifikasi".
9. Berikan saran revisi pendahuluan.
10. Buat versi pendahuluan yang lebih matang.

Aturan wajib:
- Jangan mengarang data penelitian, angka, lokasi, metode, temuan, atau konteks yang tidak ada di artikel.
- Jangan membuat kutipan palsu atau menambahkan sumber baru yang tidak tersedia di artikel.
- Jika daftar pustaka tidak ditemukan, semua status pencocokan kutipan harus ditulis "perlu verifikasi".
- Jika informasi artikel tidak cukup, tulis keterbatasannya secara eksplisit.
- Fokus pada pendahuluan, bukan mengubah seluruh artikel.
- Gunakan bahasa Indonesia akademik, jelas, formal, dan konstruktif.
- Kembalikan jawaban hanya dalam JSON valid.

Format JSON wajib:
{{
  "evaluasi_judul": "teks evaluasi",
  "rekomendasi_judul_baru": ["judul 1", "judul 2", "judul 3"],
  "evaluasi_abstrak": "teks evaluasi",
  "abstrak_versi_perbaikan": "teks abstrak perbaikan atau keterbatasan",
  "evaluasi_pendahuluan": "teks evaluasi mendalam",
  "gap_penelitian": "teks evaluasi gap",
  "cek_kutipan_dan_daftar_pustaka": [
    {{
      "kutipan": "contoh (Nama, Tahun)",
      "status": "ditemukan/perlu verifikasi",
      "catatan": "alasan singkat"
    }}
  ],
  "rumusan_masalah": "teks evaluasi rumusan masalah",
  "pendahuluan_versi_revisi": "teks pendahuluan revisi"
}}

Isi artikel:
\"\"\"{article_text}\"\"\"
""".strip()


def build_methodology_prompt(article_text, methodology_info):
    methodology_text = methodology_info["text"]
    detected_status = "terdeteksi" if methodology_info["detected"] else "tidak terdeteksi"
    detected_heading = methodology_info["heading"] or "tidak ada heading eksplisit"

    return f"""
Anda adalah reviewer metodologi artikel ilmiah berbahasa Indonesia.

Fokus Anda adalah menganalisis dan merevisi bagian METODOLOGI secara akademik, sistematis, dan jujur.

Status deteksi bagian metodologi: {detected_status}
Heading yang terdeteksi: {detected_heading}

Teks metodologi yang terdeteksi:
\"\"\"{methodology_text}\"\"\"

Tugas Anda:
1. Evaluasi apakah metodologi sudah sesuai dengan standar artikel ilmiah.
2. Periksa jenis penelitian, pendekatan penelitian, desain penelitian, subjek atau partisipan, lokasi dan waktu, populasi dan sampel, teknik pengumpulan data, instrumen, prosedur, teknik analisis data, validitas dan reliabilitas instrumen jika relevan, etika penelitian jika relevan, serta kesesuaian metode dengan rumusan masalah dan tujuan penelitian.
3. Jika metodologi kurang jelas, tidak lengkap, atau tidak sesuai, berikan catatan masalah secara rinci.
4. Revisi metodologi agar lebih sistematis, ilmiah, jelas, sesuai dengan tujuan penelitian, tidak bertele-tele, dan tidak mengarang data baru.

Aturan wajib:
- Jangan membuat data palsu.
- Jangan mengarang sampel, lokasi, waktu penelitian, instrumen, prosedur, hasil uji validitas, hasil uji reliabilitas, atau teknik analisis data yang tidak ada dalam artikel.
- Jika informasi tertentu tidak ada di artikel, tulis persis: "Informasi ini belum tersedia dalam artikel dan perlu dilengkapi penulis."
- Jangan mengubah jenis penelitian secara sembarangan. Jika metode tampak tidak konsisten, berikan saran dan alasan, bukan langsung mengganti.
- Jika bagian metodologi tidak terdeteksi, jelaskan keterbatasannya tanpa menebak isi dari section lain.
- Gunakan bahasa Indonesia akademik, jelas, formal, dan konstruktif.
- Kembalikan jawaban hanya dalam JSON valid.

Format JSON wajib:
{{
  "evaluasi_umum_metodologi": "teks evaluasi umum",
  "kelemahan_metodologi": ["kelemahan 1", "kelemahan 2"],
  "kesesuaian_metode_dengan_tujuan_rumusan_masalah": "teks evaluasi kesesuaian",
  "komponen_metodologi_yang_belum_ada": ["komponen 1", "komponen 2"],
  "saran_perbaikan": ["saran 1", "saran 2"],
  "metodologi_versi_revisi": "teks metodologi revisi",
  "catatan_yang_perlu_dilengkapi_penulis": ["catatan 1", "catatan 2"]
}}
""".strip()


def build_results_discussion_prompt(article_text, section_info, previous_result):
    section_text = section_info["text"]
    detected_status = "terdeteksi" if section_info["detected"] else "tidak terdeteksi"
    detected_heading = section_info["heading"] or "tidak ada heading eksplisit"
    previous_context = json.dumps(previous_result, ensure_ascii=False, indent=2)

    return f"""
Anda adalah reviewer akademik artikel ilmiah berbahasa Indonesia yang fokus pada bagian HASIL DAN PEMBAHASAN.

Tujuan utama Anda adalah memastikan bagian hasil dan pembahasan konsisten dengan judul penelitian, tujuan penelitian, rumusan masalah, gap penelitian, dan metodologi yang digunakan.

Status deteksi bagian hasil dan pembahasan: {detected_status}
Heading pertama yang terdeteksi: {detected_heading}

Teks hasil dan pembahasan yang terdeteksi:
\"\"\"{section_text}\"\"\"

Hasil analisis modul sebelumnya sebagai konteks pemeriksaan silang:
\"\"\"{previous_context}\"\"\"

Tugas analisis hasil penelitian:
1. Periksa apakah hasil benar-benar menjawab rumusan masalah.
2. Periksa apakah hasil sesuai dengan tujuan penelitian.
3. Periksa apakah hasil sesuai dengan metode yang digunakan.
4. Periksa apakah hasil disajikan cukup jelas.
5. Identifikasi hasil yang tidak relevan dengan tujuan penelitian.
6. Tandai klaim yang tidak didukung data.
7. Periksa apakah tabel dan grafik telah dijelaskan.
8. Periksa apakah interpretasi data sesuai dengan data yang ditampilkan.
9. Periksa inkonsistensi antara data dan narasi.
10. Identifikasi tujuan penelitian yang belum terjawab.

Tugas analisis pembahasan:
1. Periksa apakah pembahasan menjelaskan makna hasil penelitian.
2. Periksa apakah pembahasan hanya mengulang hasil tanpa analisis.
3. Periksa hubungan pembahasan dengan teori.
4. Periksa hubungan pembahasan dengan penelitian terdahulu.
5. Periksa apakah pembahasan menjelaskan persamaan dan perbedaan dengan penelitian lain.
6. Periksa apakah pembahasan menunjukkan kontribusi penelitian.
7. Periksa apakah pembahasan menjawab gap penelitian pada pendahuluan.
8. Periksa apakah pembahasan terlalu deskriptif dan kurang analitis.

Pemeriksaan silang wajib:
JUDUL -> TUJUAN PENELITIAN -> RUMUSAN MASALAH -> METODOLOGI -> HASIL -> PEMBAHASAN.
Temukan bagian yang tidak selaras, misalnya tujuan tidak seluruhnya terjawab, metode tidak cocok dengan narasi pembahasan, rumusan masalah tidak terjawab, atau hasil tidak mendukung klaim.

Aturan wajib:
- Jangan mengarang data penelitian.
- Jangan mengarang angka statistik.
- Jangan mengarang hasil observasi.
- Jangan mengarang hasil wawancara.
- Jangan mengarang tabel atau grafik.
- Jangan menambahkan kutipan atau sumber baru yang tidak tersedia dalam artikel.
- Jika data tidak tersedia, tulis persis: "Bagian ini memerlukan data tambahan dari penulis dan tidak dapat direvisi secara penuh tanpa data asli."
- Revisi hasil dan pembahasan hanya boleh merapikan, memperjelas, menata ulang, dan memperkuat argumen berdasarkan informasi yang tersedia.
- Jika hasil atau pembahasan tidak terdeteksi, jelaskan keterbatasannya dan jangan membuat isi pengganti dari nol.
- Gunakan bahasa Indonesia akademik, jelas, formal, dan konstruktif.
- Kembalikan jawaban hanya dalam JSON valid.

Format JSON wajib:
{{
  "evaluasi_hasil_penelitian": "teks evaluasi",
  "evaluasi_pembahasan": "teks evaluasi",
  "kesesuaian_dengan_rumusan_masalah": "teks evaluasi",
  "kesesuaian_dengan_metodologi": "teks evaluasi",
  "kesesuaian_dengan_tujuan_penelitian": "teks evaluasi",
  "inkonsistensi_yang_ditemukan": ["inkonsistensi 1", "inkonsistensi 2"],
  "saran_perbaikan": ["saran 1", "saran 2"],
  "hasil_versi_revisi": "teks hasil revisi atau kalimat keterbatasan",
  "pembahasan_versi_revisi": "teks pembahasan revisi atau kalimat keterbatasan",
  "catatan_untuk_penulis": ["catatan 1", "catatan 2"]
}}
""".strip()


def build_consistency_prompt(article_text, previous_result):
    previous_context = json.dumps(previous_result, ensure_ascii=False, indent=2)

    return f"""
Anda adalah reviewer jurnal ilmiah berbahasa Indonesia yang memeriksa KONSISTENSI ARTIKEL SECARA MENYELURUH.

Tujuan Anda adalah menilai apakah seluruh bagian artikel saling selaras sebagai satu naskah ilmiah utuh.

Anda wajib menggunakan dua sumber konteks:
1. Isi artikel lengkap.
2. Hasil analisis semua modul sebelumnya dalam JSON.

Hasil analisis modul sebelumnya:
\"\"\"{previous_context}\"\"\"

Bagian yang harus dibandingkan:
1. Judul
2. Abstrak
3. Pendahuluan
4. Gap penelitian
5. Tujuan penelitian
6. Rumusan masalah
7. Metodologi
8. Hasil
9. Pembahasan
10. Kesimpulan jika tersedia
11. Daftar pustaka

Pemeriksaan wajib:
1. Apakah judul sesuai dengan isi artikel.
2. Apakah abstrak sudah mencerminkan seluruh isi artikel.
3. Apakah gap penelitian di pendahuluan benar-benar dijawab oleh hasil dan pembahasan.
4. Apakah rumusan masalah dijawab oleh hasil.
5. Apakah tujuan penelitian selaras dengan metode.
6. Apakah metode cocok untuk menjawab rumusan masalah.
7. Apakah hasil sesuai dengan metode.
8. Apakah pembahasan benar-benar membahas hasil, bukan hanya mengulang.
9. Apakah kesimpulan ditarik dari hasil, bukan dari asumsi baru.
10. Apakah ada istilah, variabel, subjek, lokasi, atau fokus penelitian yang berubah-ubah antarbagian.
11. Apakah ada klaim besar yang tidak didukung data.
12. Apakah ada data di hasil yang tidak dibahas.
13. Apakah ada pembahasan yang tidak memiliki dasar data.
14. Apakah kutipan penting sesuai dengan daftar pustaka.
15. Apakah artikel sudah tampak utuh sebagai satu naskah ilmiah.

Aturan wajib:
- Jangan mengarang data.
- Jangan menambah hasil penelitian baru.
- Jangan membuat kutipan palsu.
- Jangan menambah sumber baru.
- Jangan mengubah substansi artikel melampaui isi yang tersedia.
- Revisi dan rekomendasi hanya boleh berdasarkan isi artikel dan hasil analisis modul sebelumnya.
- Jika ada bagian yang tidak tersedia, tulis persis: "Bagian ini belum tersedia dalam artikel sehingga konsistensinya belum dapat dinilai secara penuh."
- Gunakan bahasa Indonesia akademik, jelas, formal, dan konstruktif.
- Kembalikan jawaban hanya dalam JSON valid.

Format JSON wajib:
{{
  "ringkasan_kualitas_konsistensi_artikel": "ringkasan kualitas konsistensi artikel secara menyeluruh",
  "skor_konsistensi_artikel": 75,
  "peta_keselarasan_artikel": {{
    "judul_abstrak": "evaluasi Judul <-> Abstrak",
    "pendahuluan_metodologi": "evaluasi Pendahuluan <-> Metodologi",
    "rumusan_masalah_hasil": "evaluasi Rumusan Masalah <-> Hasil",
    "hasil_pembahasan": "evaluasi Hasil <-> Pembahasan",
    "pembahasan_kesimpulan": "evaluasi Pembahasan <-> Kesimpulan",
    "kutipan_daftar_pustaka": "evaluasi Kutipan <-> Daftar Pustaka"
  }},
  "daftar_inkonsistensi_yang_ditemukan": ["inkonsistensi 1", "inkonsistensi 2"],
  "dampak_akademik_dari_inkonsistensi": ["dampak 1", "dampak 2"],
  "prioritas_perbaikan": {{
    "sangat_penting": ["prioritas 1"],
    "penting": ["prioritas 2"],
    "tambahan": ["prioritas 3"]
  }},
  "rekomendasi_revisi_menyeluruh": ["rekomendasi 1", "rekomendasi 2"],
  "checklist_kesiapan_artikel_untuk_submit_jurnal": [
    {{
      "aspek": "Judul dan abstrak",
      "status": "siap/perlu revisi/belum dapat dinilai",
      "catatan": "catatan singkat"
    }}
  ]
}}

Isi artikel lengkap:
\"\"\"{article_text}\"\"\"
""".strip()


def build_reviewer_simulator_prompt(article_text, previous_result):
    previous_context = json.dumps(previous_result, ensure_ascii=False, indent=2)

    return f"""
Anda adalah reviewer jurnal bereputasi yang melakukan peer review artikel ilmiah secara ketat, adil, dan konstruktif.

Anda harus membaca dan menilai:
1. Judul
2. Abstrak
3. Pendahuluan
4. Metodologi
5. Hasil dan Pembahasan
6. Kesimpulan
7. Daftar Pustaka

Gunakan isi artikel lengkap dan hasil analisis modul sebelumnya sebagai konteks. Jangan hanya memberi komentar umum; tulis alasan akademik yang rinci seperti reviewer jurnal.

Hasil analisis modul sebelumnya:
\"\"\"{previous_context}\"\"\"

Aturan wajib:
- Jangan mengarang data.
- Jangan mengarang hasil penelitian.
- Jangan membuat kutipan palsu.
- Jangan menambahkan referensi baru yang tidak ada dalam artikel.
- Jika bagian tertentu belum tersedia, jelaskan bahwa penilaian bagian tersebut terbatas.
- Keputusan reviewer harus salah satu dari: Accept, Minor Revision, Major Revision, Reject.
- Gunakan bahasa Indonesia akademik, jelas, formal, tajam, dan konstruktif.
- Kembalikan jawaban hanya dalam JSON valid.

Format JSON wajib:
{{
  "kelebihan_artikel": ["kelebihan 1", "kelebihan 2"],
  "kelemahan_artikel": ["kelemahan 1", "kelemahan 2"],
  "novelty_penelitian": "evaluasi novelty penelitian",
  "signifikansi_penelitian": "evaluasi signifikansi penelitian",
  "kekuatan_metodologi": ["kekuatan 1", "kekuatan 2"],
  "kelemahan_metodologi": ["kelemahan 1", "kelemahan 2"],
  "kualitas_pembahasan": "evaluasi kualitas pembahasan",
  "kualitas_referensi": "evaluasi kualitas referensi",
  "kualitas_bahasa_akademik": "evaluasi kualitas bahasa akademik",
  "potensi_publikasi": "evaluasi potensi publikasi",
  "keputusan_reviewer": "Accept/Minor Revision/Major Revision/Reject",
  "alasan_keputusan": "alasan rinci keputusan reviewer"
}}

Isi artikel lengkap:
\"\"\"{article_text}\"\"\"
""".strip()


def build_reference_citation_prompt(article_text, previous_result):
    previous_context = json.dumps(previous_result, ensure_ascii=False, indent=2)

    return f"""
Anda adalah reviewer akademik yang memeriksa kualitas REFERENSI DAN SITASI artikel ilmiah secara menyeluruh.

Tanggal penilaian: 29 Mei 2026. Gunakan tanggal ini untuk menghitung usia referensi.

Tujuan:
Evaluasi apakah sitasi dan daftar pustaka sudah memenuhi standar artikel ilmiah yang layak dipublikasikan.

Gunakan isi artikel lengkap dan hasil analisis modul sebelumnya sebagai konteks. Jangan hanya menilai daftar pustaka; cocokkan juga sitasi dalam pendahuluan, metodologi, hasil dan pembahasan, serta kesimpulan jika tersedia.

Hasil analisis modul sebelumnya:
\"\"\"{previous_context}\"\"\"

Pemeriksaan sitasi wajib:
1. Periksa semua sitasi dalam naskah.
2. Identifikasi sitasi yang tidak memiliki pasangan di daftar pustaka.
3. Identifikasi nama penulis yang berbeda dengan daftar pustaka.
4. Identifikasi tahun yang tidak konsisten.
5. Identifikasi sitasi yang tidak lengkap.
6. Identifikasi sitasi yang diduga salah format.
7. Identifikasi sitasi yang berulang secara berlebihan.
8. Identifikasi self-citation berlebihan jika terdeteksi dari artikel.

Pemeriksaan daftar pustaka wajib:
1. Referensi yang tidak pernah disitasi.
2. Referensi yang hanya muncul di daftar pustaka.
3. Referensi duplikat.
4. Referensi yang tidak lengkap.
5. Referensi yang diduga salah format.
6. Konsistensi format seluruh daftar pustaka.

Pemeriksaan kualitas referensi wajib:
1. Proporsi jurnal ilmiah.
2. Proporsi buku.
3. Proporsi prosiding.
4. Proporsi website.
5. Proporsi sumber primer.
6. Proporsi sumber sekunder.
7. Keragaman referensi.
8. Dominasi referensi dari satu penulis.
9. Dominasi referensi dari satu tahun.
10. Keseimbangan referensi nasional dan internasional.

Pemeriksaan usia referensi wajib:
- Hitung referensi <= 5 tahun.
- Hitung referensi 6-10 tahun.
- Hitung referensi > 10 tahun.
- Nilai apakah referensi cukup mutakhir untuk bidang penelitian.
- Jika banyak referensi lama, berikan peringatan.

Pemeriksaan format wajib:
- Bandingkan format referensi dengan APA 7th Edition.
- Jika ada kesalahan format, tampilkan referensi asli dan contoh format yang benar hanya berdasarkan informasi bibliografi yang tersedia.

Validasi akademik wajib:
- Identifikasi referensi yang tampak tidak ilmiah.
- Identifikasi referensi yang kurang kredibel.
- Identifikasi referensi yang sebaiknya diganti.
- Identifikasi referensi yang perlu diperbarui.

Aturan wajib:
- Jangan mengarang DOI.
- Jangan mengarang URL.
- Jangan mengarang data bibliografi.
- Jangan menambahkan penulis, judul, nama jurnal, penerbit, volume, nomor, halaman, DOI, atau URL yang tidak tersedia.
- Jika keberadaan atau kredibilitas referensi tidak dapat dipastikan dari artikel saja, tulis persis: "Referensi perlu diverifikasi pada sumber asli."
- Jika daftar pustaka tidak tersedia, jelaskan bahwa penilaian daftar pustaka tidak dapat dilakukan secara penuh.
- Gunakan bahasa Indonesia akademik, jelas, formal, dan konstruktif.
- Kembalikan jawaban hanya dalam JSON valid.

Format JSON wajib:
{{
  "ringkasan_kualitas_referensi": "ringkasan penilaian",
  "statistik_referensi": {{
    "jumlah_total_referensi": 0,
    "jurnal": 0,
    "buku": 0,
    "prosiding": 0,
    "website": 0,
    "lainnya": 0,
    "catatan": "catatan jika klasifikasi tidak pasti"
  }},
  "statistik_usia_referensi": {{
    "referensi_5_tahun_atau_kurang": 0,
    "referensi_6_10_tahun": 0,
    "referensi_lebih_dari_10_tahun": 0,
    "evaluasi_kemutakhiran": "evaluasi"
  }},
  "sitasi_bermasalah": [
    {{
      "sitasi": "sitasi dalam naskah",
      "masalah": "jenis masalah",
      "catatan": "catatan"
    }}
  ],
  "referensi_bermasalah": [
    {{
      "referensi": "referensi asli",
      "masalah": "jenis masalah",
      "catatan": "Referensi perlu diverifikasi pada sumber asli."
    }}
  ],
  "referensi_tidak_pernah_disitasi": ["referensi 1"],
  "sitasi_tanpa_referensi": ["sitasi 1"],
  "referensi_duplikat": ["duplikasi 1"],
  "evaluasi_kualitas_referensi": "evaluasi kualitas referensi",
  "evaluasi_kepatuhan_apa_7": [
    {{
      "referensi_asli": "referensi asli",
      "masalah_format": "masalah",
      "contoh_format_apa_7": "contoh berdasarkan data yang tersedia"
    }}
  ],
  "daftar_perbaikan_yang_direkomendasikan": ["perbaikan 1", "perbaikan 2"],
  "skor_kualitas_referensi": 75,
  "tingkat_kesiapan_referensi_untuk_publikasi": "Kurang Layak/Cukup/Baik/Sangat Baik/Siap Publikasi"
}}

Isi artikel lengkap:
\"\"\"{article_text}\"\"\"
""".strip()


def build_novelty_contribution_prompt(article_text, previous_result):
    previous_context = json.dumps(previous_result, ensure_ascii=False, indent=2)

    return f"""
Anda adalah reviewer akademik yang mengevaluasi NOVELTY, kontribusi ilmiah, dan nilai tambah penelitian berdasarkan keseluruhan artikel.

Gunakan konteks dari judul, abstrak, pendahuluan, gap penelitian, rumusan masalah, metodologi, hasil, pembahasan, kesimpulan, daftar pustaka, serta hasil analisis modul sebelumnya.

Hasil analisis modul sebelumnya:
\"\"\"{previous_context}\"\"\"

Tujuan penilaian:
1. Tentukan apa yang baru dari penelitian ini.
2. Tentukan kontribusi penelitian ini.
3. Nilai seberapa kuat novelty penelitian.
4. Nilai peluang artikel dianggap memiliki kontribusi ilmiah.

Identifikasi novelty:
- Periksa apakah penelitian menawarkan teori baru, model baru, framework baru, metode baru, instrumen baru, media baru, aplikasi baru, pendekatan baru, konteks baru, populasi baru, atau lokasi penelitian yang unik.
- Tentukan apakah novelty dinyatakan eksplisit, hanya tersirat, belum terlihat, atau hanya mengulang penelitian terdahulu.

Analisis gap penelitian:
1. Apakah gap pada pendahuluan nyata dan jelas.
2. Apakah gap dijawab oleh penelitian.
3. Apakah hasil berhasil mengisi gap tersebut.
4. Apakah ada ketidaksesuaian antara gap dan hasil penelitian.

Analisis kontribusi ilmiah:
A. Kontribusi Teoretis: pengembangan teori, penguatan teori, pengujian teori, atau perluasan teori.
B. Kontribusi Praktis: manfaat bagi sekolah, guru, siswa, institusi, masyarakat, pemerintah, industri, atau pemangku kepentingan lain yang memang disebut dalam artikel.
C. Kontribusi Metodologis: metode baru, instrumen baru, teknik analisis baru, atau desain penelitian baru.

Klasifikasikan tingkat kebaruan:
- Level 1: Replikasi murni.
- Level 2: Replikasi dengan konteks berbeda.
- Level 3: Modifikasi kecil.
- Level 4: Pengembangan moderat.
- Level 5: Kontribusi signifikan.
- Level 6: Novelty tinggi.
- Level 7: Potensi kontribusi besar terhadap bidang ilmu.

Identifikasi risiko akademik:
- penelitian terlalu umum;
- penelitian terlalu sering dilakukan;
- novelty tidak terlihat;
- kontribusi terlalu kecil;
- hasil tidak cukup untuk mendukung klaim kontribusi;
- gap penelitian lemah.

Berikan saran:
1. Cara memperjelas novelty.
2. Cara memperkuat kontribusi.
3. Cara menulis novelty pada pendahuluan.
4. Cara menampilkan kontribusi pada pembahasan.
5. Cara memperkuat posisi artikel terhadap penelitian terdahulu.

Aturan wajib:
- Jangan mengarang penelitian terdahulu.
- Jangan mengarang teori.
- Jangan mengarang hasil penelitian.
- Jangan mengarang data baru.
- Jangan menambahkan kutipan atau referensi baru.
- Jika informasi tidak cukup untuk menilai novelty secara penuh, tulis persis: "Novelty memerlukan verifikasi lebih lanjut melalui penelusuran literatur yang lebih luas."
- Gunakan bahasa Indonesia akademik, jelas, formal, dan konstruktif.
- Kembalikan jawaban hanya dalam JSON valid.

Format JSON wajib:
{{
  "ringkasan_novelty_penelitian": "ringkasan novelty",
  "ringkasan_kontribusi_ilmiah": "ringkasan kontribusi",
  "evaluasi_gap_penelitian": "evaluasi gap",
  "novelty_yang_teridentifikasi": ["novelty 1", "novelty 2"],
  "kontribusi_teoretis": "evaluasi kontribusi teoretis",
  "kontribusi_praktis": "evaluasi kontribusi praktis",
  "kontribusi_metodologis": "evaluasi kontribusi metodologis",
  "risiko_akademik": ["risiko 1", "risiko 2"],
  "saran_penguatan_novelty": ["saran 1", "saran 2"],
  "saran_penguatan_kontribusi": ["saran 1", "saran 2"],
  "skor_novelty": 70,
  "skor_kontribusi_ilmiah": 75,
  "tingkat_kebaruan_penelitian": {{
    "level": "Level 4",
    "kategori": "Pengembangan moderat",
    "alasan": "alasan klasifikasi"
  }},
  "kelayakan_kontribusi_untuk_publikasi": "Rendah/Cukup/Baik/Sangat Baik/Unggul"
}}

Isi artikel lengkap:
\"\"\"{article_text}\"\"\"
""".strip()


def build_submission_readiness_prompt(article_text, previous_result):
    previous_context = json.dumps(previous_result, ensure_ascii=False, indent=2)

    return f"""
Anda adalah editor jurnal yang melakukan pemeriksaan akhir sebelum artikel disubmit.

Tugas Anda adalah menggabungkan seluruh hasil analisis sebelumnya untuk menilai apakah artikel sudah siap dikirim ke jurnal.

Gunakan seluruh konteks yang tersedia:
- Judul
- Abstrak
- Kata kunci
- Pendahuluan
- Gap penelitian
- Rumusan masalah
- Metodologi
- Hasil
- Pembahasan
- Kesimpulan
- Referensi
- Sitasi
- Novelty
- Kontribusi ilmiah
- Konsistensi artikel

Hasil analisis semua modul sebelumnya:
\"\"\"{previous_context}\"\"\"

Pemeriksaan struktur artikel:
Periksa kelengkapan dan kelayakan judul, abstrak, kata kunci, pendahuluan, metodologi, hasil, pembahasan, kesimpulan, dan daftar pustaka. Tandai bagian yang hilang atau belum memadai.

Pemeriksaan kualitas ilmiah:
1. Kejelasan masalah penelitian.
2. Kekuatan gap penelitian.
3. Kejelasan tujuan penelitian.
4. Kesesuaian metodologi.
5. Kualitas hasil penelitian.
6. Kualitas pembahasan.
7. Kualitas kesimpulan.
8. Kualitas referensi.
9. Kualitas argumentasi ilmiah.
10. Kualitas kontribusi penelitian.

Pemeriksaan konsistensi:
Periksa keselarasan Judul -> Abstrak -> Pendahuluan -> Rumusan Masalah -> Metodologi -> Hasil -> Pembahasan -> Kesimpulan. Identifikasi bagian yang tidak konsisten.

Pemeriksaan referensi:
Periksa kelengkapan sitasi, kualitas referensi, kesesuaian daftar pustaka, kepatuhan APA 7, dan kemutakhiran referensi.

Pemeriksaan novelty:
Periksa tingkat kebaruan penelitian, kontribusi teoretis, kontribusi praktis, dan kontribusi metodologis.

Simulasi reviewer:
Perkirakan bagaimana reviewer jurnal kemungkinan menilai artikel ini. Tampilkan kekuatan utama, kelemahan utama, potensi masalah saat review, dan bagian yang berisiko dikritik reviewer.

Checklist submit wajib:
- Judul Layak
- Abstrak Layak
- Pendahuluan Layak
- Gap Penelitian Jelas
- Rumusan Masalah Jelas
- Metodologi Memadai
- Hasil Memadai
- Pembahasan Memadai
- Kesimpulan Memadai
- Referensi Memadai
- Novelty Terlihat
- Kontribusi Jelas
- Konsistensi Terjaga

Hitung skor:
1. Skor Struktur Artikel (0-100)
2. Skor Kualitas Ilmiah (0-100)
3. Skor Referensi (0-100)
4. Skor Novelty (0-100)
5. Skor Konsistensi (0-100)
6. Skor Kesiapan Submit Keseluruhan (0-100)

Tentukan satu kategori kelayakan:
- Belum Layak Submit
- Perlu Revisi Besar
- Perlu Revisi Minor
- Siap Submit Nasional
- Siap Submit SINTA
- Siap Submit Internasional
- Sangat Siap Submit

Rencana aksi:
Berikan perbaikan prioritas tinggi, sedang, dan rendah, urutkan berdasarkan dampaknya terhadap peluang diterima.

Aturan wajib:
- Jangan mengarang data penelitian.
- Jangan mengarang referensi.
- Jangan mengarang hasil penelitian.
- Jangan mengarang novelty yang tidak ada.
- Jika informasi tidak cukup, tulis persis: "Penilaian memerlukan informasi tambahan dari artikel."
- Gunakan bahasa Indonesia akademik, jelas, formal, dan konstruktif.
- Kembalikan jawaban hanya dalam JSON valid.

Format JSON wajib:
{{
  "ringkasan_evaluasi_artikel": "ringkasan akhir editor",
  "checklist_kesiapan_submit": [
    {{
      "aspek": "Judul Layak",
      "status": "layak/perlu revisi/belum memadai/belum dapat dinilai",
      "catatan": "catatan singkat"
    }}
  ],
  "skor_kesiapan_submit": {{
    "struktur_artikel": 70,
    "kualitas_ilmiah": 70,
    "referensi": 70,
    "novelty": 70,
    "konsistensi": 70,
    "keseluruhan": 70
  }},
  "kategori_kelayakan": "Belum Layak Submit/Perlu Revisi Besar/Perlu Revisi Minor/Siap Submit Nasional/Siap Submit SINTA/Siap Submit Internasional/Sangat Siap Submit",
  "kelebihan_artikel": ["kelebihan 1", "kelebihan 2"],
  "kelemahan_artikel": ["kelemahan 1", "kelemahan 2"],
  "risiko_penolakan": ["risiko 1", "risiko 2"],
  "prioritas_revisi": {{
    "prioritas_tinggi": ["perbaikan 1"],
    "prioritas_sedang": ["perbaikan 2"],
    "prioritas_rendah": ["perbaikan 3"]
  }},
  "rekomendasi_akhir_editor": "rekomendasi akhir"
}}

Isi artikel lengkap:
\"\"\"{article_text}\"\"\"
""".strip()


def call_openai_chat(system_prompt, user_prompt):
    response = client.chat.completions.create(
        model=os.getenv(
            "OPENAI_MODEL",
            "gpt-5.6",
        ),
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
    )

    content = response.choices[0].message.content
    if not content:
        raise RuntimeError("OpenAI API tidak mengembalikan konten analisis.")

    return content


def openai_chat_json(system_prompt, user_prompt):
    content = call_openai_chat(system_prompt, user_prompt).strip()
    if content.startswith("```") and content.endswith("```"):
        content = re.sub(r"^```(?:json)?\s*", "", content, count=1)
        content = re.sub(r"\s*```$", "", content, count=1)

    try:
        return json.loads(content)
    except json.JSONDecodeError as exc:
        raise RuntimeError("OpenAI API mengembalikan respons yang bukan JSON valid. Coba ulangi analisis.") from exc


def test_openai_connection():
    return call_openai_chat(
        "Jawab secara singkat dan jelas dalam bahasa Indonesia.",
        "Halo",
    )


REVISION_SYSTEM_PROMPT = (
    "Anda adalah editor artikel ilmiah. Berikan revisi final yang ringkas dan siap dipakai. "
    "Jangan mengarang data, angka, lokasi, metode, hasil, sitasi, atau referensi baru. "
    "Jika informasi tidak tersedia, nyatakan bahwa bagian tersebut perlu dilengkapi penulis. "
    "Kembalikan hanya JSON valid sesuai format."
)


def analyze_title_with_ai(title_text, article_context):
    prompt = f"""
Revisi judul artikel berdasarkan judul asli dan konteks singkat.

Judul asli:
\"\"\"{title_text}\"\"\"

Konteks singkat artikel:
\"\"\"{article_context}\"\"\"

Format JSON:
{{
  "judul_awal": "...",
  "alternatif_judul": ["...", "...", "..."],
  "judul_rekomendasi": "..."
}}
""".strip()
    return openai_chat_json(REVISION_SYSTEM_PROMPT, prompt)


def analyze_abstract_with_ai(abstract_text, title_result, article_context):
    prompt = f"""
Revisi abstrak agar selaras dengan judul rekomendasi. Pertahankan seluruh fakta yang tersedia.

Judul rekomendasi: {title_result.get("judul_rekomendasi", "")}
Konteks singkat:
\"\"\"{article_context}\"\"\"
Abstrak asli:
\"\"\"{abstract_text}\"\"\"

Format JSON:
{{"abstrak_awal": "...", "abstrak_revisi": "..."}}
""".strip()
    return openai_chat_json(REVISION_SYSTEM_PROMPT, prompt)


def analyze_introduction_with_ai(introduction_text, title_result, abstract_result):
    prompt = f"""

Revisi pendahuluan agar koheren dengan judul rekomendasi dan abstrak revisi.

Perbaiki:
- alur argumentasi
- latar belakang
- gap penelitian
- tujuan penelitian

Pertahankan sitasi yang ada.
Jangan menambah sitasi baru.
Jangan mengarang data baru.

Tampilkan hanya pendahuluan hasil revisi.

Judul rekomendasi: {title_result.get("judul_rekomendasi", "")}
Abstrak revisi:
\"\"\"{abstract_result.get("abstrak_revisi", "")}\"\"\"
Pendahuluan asli:
\"\"\"{introduction_text}\"\"\"

Format JSON:
{{"pendahuluan_awal": "...", "pendahuluan_revisi": "..."}}
""".strip()
    return openai_chat_json(REVISION_SYSTEM_PROMPT, prompt)


def analyze_methodology_with_ai(methodology_text, title_result, abstract_result, introduction_result):
    prompt = f"""
Revisi metodologi agar jelas dan konsisten dengan tahap sebelumnya. Jangan menambahkan detail penelitian yang tidak tersedia.

Judul rekomendasi: {title_result.get("judul_rekomendasi", "")}
Abstrak revisi:
\"\"\"{abstract_result.get("abstrak_revisi", "")}\"\"\"
Pendahuluan revisi:
\"\"\"{introduction_result.get("pendahuluan_revisi", "")}\"\"\"
Metodologi asli:
\"\"\"{methodology_text}\"\"\"

Format JSON:
{{"metodologi_awal": "...", "metodologi_revisi": "..."}}
""".strip()
    return openai_chat_json(REVISION_SYSTEM_PROMPT, prompt)


def analyze_results_discussion_with_ai(
    section_text,
    title_result,
    abstract_result,
    introduction_result,
    methodology_result,
):
    prompt = f"""
Revisi hasil dan pembahasan agar koheren dengan bagian sebelumnya. Jangan mengubah atau menciptakan hasil penelitian.

Judul rekomendasi: {title_result.get("judul_rekomendasi", "")}
Abstrak revisi:
\"\"\"{abstract_result.get("abstrak_revisi", "")}\"\"\"
Pendahuluan revisi:
\"\"\"{introduction_result.get("pendahuluan_revisi", "")}\"\"\"
Metodologi revisi:
\"\"\"{methodology_result.get("metodologi_revisi", "")}\"\"\"
Hasil dan pembahasan asli:
\"\"\"{section_text}\"\"\"

Format JSON:
{{"hasil_pembahasan_awal": "...", "hasil_pembahasan_revisi": "..."}}
""".strip()
    return openai_chat_json(REVISION_SYSTEM_PROMPT, prompt)


def analyze_conclusion_with_ai(conclusion_text, result):
    prompt = f"""
Revisi kesimpulan agar hanya merangkum temuan yang tersedia dan selaras dengan seluruh revisi sebelumnya.

Judul rekomendasi: {result["judul"].get("judul_rekomendasi", "")}
Abstrak revisi:
\"\"\"{result["abstrak"].get("abstrak_revisi", "")}\"\"\"
Pendahuluan revisi:
\"\"\"{result["pendahuluan"].get("pendahuluan_revisi", "")}\"\"\"
Metodologi revisi:
\"\"\"{result["metodologi"].get("metodologi_revisi", "")}\"\"\"
Hasil dan pembahasan revisi:
\"\"\"{result["hasil_pembahasan"].get("hasil_pembahasan_revisi", "")}\"\"\"
Kesimpulan asli:
\"\"\"{conclusion_text}\"\"\"

Format JSON:
{{"kesimpulan_awal": "...", "kesimpulan_revisi": "..."}}
""".strip()
    return openai_chat_json(REVISION_SYSTEM_PROMPT, prompt)


def analyze_reference_citation_with_ai(article_text, reference_text, result):
    revision_context = {
        "judul_rekomendasi": result["judul"].get("judul_rekomendasi", ""),
        "abstrak_revisi": result["abstrak"].get("abstrak_revisi", ""),
        "pendahuluan_revisi": result["pendahuluan"].get("pendahuluan_revisi", ""),
        "metodologi_revisi": result["metodologi"].get("metodologi_revisi", ""),
        "hasil_pembahasan_revisi": result["hasil_pembahasan"].get("hasil_pembahasan_revisi", ""),
        "kesimpulan_revisi": result["kesimpulan"].get("kesimpulan_revisi", ""),
    }
    prompt = f"""
Periksa kecocokan sitasi dengan daftar pustaka dan format APA 7. Jangan mengarang DOI, URL, atau data bibliografi.

Artikel lengkap:
\"\"\"{article_text}\"\"\"

Daftar pustaka terdeteksi:
\"\"\"{reference_text}\"\"\"

Konteks hasil revisi:
{json.dumps(revision_context, ensure_ascii=False)}

Format JSON:
{{
  "sitasi_tanpa_referensi": [{{"sitasi": "...", "catatan": "..."}}],
  "referensi_tidak_disitasi": ["..."],
  "referensi_bermasalah": [
    {{"referensi_asli": "...", "masalah": "...", "revisi_apa_7": "..."}}
  ],
  "catatan_akhir_sitasi": "..."
}}
""".strip()
    return openai_chat_json(REVISION_SYSTEM_PROMPT, prompt)


def extract_reference_search_terms(result):
    prompt = f"""
Ekstrak kata kunci pencarian artikel ilmiah dari hasil revisi berikut.
Berikan 5-8 kata kunci dan 3 query pencarian singkat. Jangan menambahkan topik yang tidak ada.

Judul:
{result["judul"].get("judul_rekomendasi", "")}

Abstrak:
{result["abstrak"].get("abstrak_revisi", "")}

Pendahuluan:
{result["pendahuluan"].get("pendahuluan_revisi", "")}

Metodologi:
{result["metodologi"].get("metodologi_revisi", "")}

Format JSON:
{{"kata_kunci": ["..."], "query_pencarian": ["...", "...", "..."]}}
""".strip()
    data = openai_chat_json(REVISION_SYSTEM_PROMPT, prompt)
    keywords = [str(item).strip() for item in data.get("kata_kunci", []) if str(item).strip()]
    queries = [str(item).strip() for item in data.get("query_pencarian", []) if str(item).strip()]
    title = result["judul"].get("judul_rekomendasi", "").strip()
    if title:
        queries.insert(0, title)
    return {
        "kata_kunci": keywords[:8],
        "query_pencarian": list(dict.fromkeys(queries))[:4],
    }


def normalize_doi(value):
    if not value:
        return ""
    doi = str(value).strip().lower()
    doi = re.sub(r"^https?://(?:dx\.)?doi\.org/", "", doi)
    return doi.strip()


def format_authors_apa(authors):
    def format_name(name):
        value = str(name).strip()
        if not value:
            return ""
        if "," in value:
            family, given = [part.strip() for part in value.split(",", 1)]
        else:
            parts = value.split()
            if len(parts) == 1:
                return parts[0]
            family, given = parts[-1], " ".join(parts[:-1])
        initials = " ".join(
            f"{part[0].upper()}."
            for part in re.findall(r"[A-Za-zÀ-ÖØ-öø-ÿ]+", given)
            if part
        )
        return f"{family}, {initials}".strip().rstrip(",")

    names = [format_name(name) for name in authors if str(name).strip()]
    if not names:
        return ""
    if len(names) == 1:
        return names[0]
    if len(names) == 2:
        return f"{names[0]}, & {names[1]}"
    return f"{', '.join(names[:-1])}, & {names[-1]}"


def format_reference_apa_7(reference):
    authors = format_authors_apa(reference.get("penulis_list", []))
    year = reference.get("tahun") or "n.d."
    title = reference.get("judul", "").rstrip(".")
    journal = reference.get("jurnal", "").rstrip(".")
    doi = normalize_doi(reference.get("doi"))
    prefix = f"{authors} ({year})." if authors else f"({year})."
    parts = [prefix, f"{title}."]
    if journal:
        parts.append(f"{journal}.")
    if doi:
        parts.append(f"https://doi.org/{doi}")
    return " ".join(parts)


def crossref_reference_from_item(item):
    title_values = item.get("title") or []
    container_values = item.get("container-title") or []
    authors = []
    for author in item.get("author") or []:
        family = str(author.get("family") or "").strip()
        given = str(author.get("given") or "").strip()
        name = ", ".join(part for part in (family, given) if part)
        if name:
            authors.append(name)
    date_parts = (
        item.get("published-print", {}).get("date-parts")
        or item.get("published-online", {}).get("date-parts")
        or item.get("issued", {}).get("date-parts")
        or []
    )
    year = date_parts[0][0] if date_parts and date_parts[0] else ""
    reference = {
        "judul": str(title_values[0]).strip() if title_values else "",
        "penulis": "; ".join(authors),
        "penulis_list": authors,
        "tahun": str(year),
        "jurnal": str(container_values[0]).strip() if container_values else "",
        "quartile_atau_sinta": "Perlu verifikasi pada Scopus/WoS/SINTA",
        "doi": normalize_doi(item.get("DOI")),
        "sumber_metadata": "Crossref",
        "jumlah_sitasi": int(item.get("is-referenced-by-count") or 0),
    }
    reference["apa_7"] = format_reference_apa_7(reference)
    return reference


def openalex_reference_from_item(item):
    authors = [
        authorship.get("author", {}).get("display_name", "").strip()
        for authorship in item.get("authorships") or []
        if authorship.get("author", {}).get("display_name", "").strip()
    ]
    primary_location = item.get("primary_location") or {}
    source = primary_location.get("source") or {}
    reference = {
        "judul": str(item.get("display_name") or item.get("title") or "").strip(),
        "penulis": "; ".join(authors),
        "penulis_list": authors,
        "tahun": str(item.get("publication_year") or ""),
        "jurnal": str(source.get("display_name") or "").strip(),
        "quartile_atau_sinta": "Perlu verifikasi pada Scopus/WoS/SINTA",
        "doi": normalize_doi((item.get("ids") or {}).get("doi") or item.get("doi")),
        "sumber_metadata": "OpenAlex",
        "jumlah_sitasi": int(item.get("cited_by_count") or 0),
    }
    reference["apa_7"] = format_reference_apa_7(reference)
    return reference


def reference_is_complete(reference):
    return all(
        reference.get(key)
        for key in ("judul", "penulis", "tahun", "jurnal", "doi")
    )


def search_crossref(query, limit=15):
    params = {
        "query.bibliographic": query,
        "filter": "type:journal-article",
        "rows": limit,
    }
    email = os.getenv("SCHOLARLY_API_EMAIL", "").strip()
    if email:
        params["mailto"] = email
    response = httpx.get(
        "https://api.crossref.org/works",
        params=params,
        headers={"User-Agent": "RevisiArtikelAI/1.0"},
        timeout=30,
    )
    response.raise_for_status()
    items = response.json().get("message", {}).get("items", [])
    return [crossref_reference_from_item(item) for item in items]


def search_openalex(query, limit=15):
    params = {
        "search": query,
        "filter": "type:article,has_doi:true",
        "per-page": limit,
    }
    email = os.getenv("SCHOLARLY_API_EMAIL", "").strip()
    if email:
        params["mailto"] = email
    response = httpx.get(
        "https://api.openalex.org/works",
        params=params,
        headers={"User-Agent": "RevisiArtikelAI/1.0"},
        timeout=30,
    )
    response.raise_for_status()
    return [
        openalex_reference_from_item(item)
        for item in response.json().get("results", [])
    ]


def rank_reference_candidates(candidates, result, search_terms):
    immutable_candidates = [
        {
            "id": index,
            "judul": item["judul"],
            "tahun": item["tahun"],
            "jurnal": item["jurnal"],
        }
        for index, item in enumerate(candidates)
    ]
    prompt = f"""
Urutkan kandidat referensi nyata berikut berdasarkan relevansi dengan artikel.
Jangan mengubah metadata dan jangan menambahkan referensi. Gunakan hanya ID yang tersedia.

Konteks artikel:
Judul: {result["judul"].get("judul_rekomendasi", "")}
Abstrak: {result["abstrak"].get("abstrak_revisi", "")}
Pendahuluan: {result["pendahuluan"].get("pendahuluan_revisi", "")}
Metodologi: {result["metodologi"].get("metodologi_revisi", "")}
Kata kunci: {", ".join(search_terms.get("kata_kunci", []))}

Kandidat:
{json.dumps(immutable_candidates, ensure_ascii=False)}

Kategori yang diizinkan:
1. Sangat relevan dengan topik
2. Sangat relevan dengan metodologi
3. Pendukung teori dan pembahasan

Format JSON:
{{
  "peringkat": [
    {{"id": 0, "kategori": "1", "alasan_relevan": "alasan singkat berbasis judul kandidat"}}
  ]
}}
""".strip()
    try:
        ranked = openai_chat_json(REVISION_SYSTEM_PROMPT, prompt).get("peringkat", [])
    except Exception:
        ranked = []

    ranking_map = {}
    for position, item in enumerate(ranked):
        try:
            candidate_id = int(item.get("id"))
        except (TypeError, ValueError):
            continue
        if 0 <= candidate_id < len(candidates) and candidate_id not in ranking_map:
            category = str(item.get("kategori") or "3").strip()
            category = category[0] if category and category[0] in {"1", "2", "3"} else "3"
            ranking_map[candidate_id] = {
                "position": position,
                "kategori": category,
                "alasan_relevan": str(item.get("alasan_relevan") or "").strip(),
            }

    ordered = []
    for index, reference in enumerate(candidates):
        annotation = ranking_map.get(index, {})
        item = dict(reference)
        item["kategori"] = annotation.get("kategori", "3")
        item["alasan_relevan"] = annotation.get(
            "alasan_relevan",
            "Relevan berdasarkan kemiripan metadata dengan kata kunci artikel.",
        )
        item["_rank"] = annotation.get("position", len(candidates) + index)
        ordered.append(item)
    ordered.sort(key=lambda item: (item["_rank"], -item.get("jumlah_sitasi", 0)))
    for item in ordered:
        item.pop("_rank", None)
    return ordered


def find_verified_reference_recommendations(result):
    search_terms = extract_reference_search_terms(result)
    candidates_by_doi = {}
    errors = []

    search_jobs = [
        (search_function, query)
        for query in search_terms["query_pencarian"]
        for search_function in (search_openalex, search_crossref)
    ]
    if not search_jobs:
        raise RuntimeError("Kata kunci pencarian referensi tidak berhasil diekstrak.")
    with ThreadPoolExecutor(max_workers=min(8, len(search_jobs))) as executor:
        futures = {
            executor.submit(search_function, query): (search_function, query)
            for search_function, query in search_jobs
        }
        for future in as_completed(futures):
            search_function, query = futures[future]
            try:
                references = future.result()
            except Exception as exc:
                errors.append(f"{search_function.__name__} ({query}): {exc}")
                continue
            for reference in references:
                doi = normalize_doi(reference.get("doi"))
                if doi and reference_is_complete(reference):
                    existing = candidates_by_doi.get(doi)
                    if not existing or reference.get("jumlah_sitasi", 0) > existing.get("jumlah_sitasi", 0):
                        candidates_by_doi[doi] = reference

    candidates = sorted(
        candidates_by_doi.values(),
        key=lambda item: item.get("jumlah_sitasi", 0),
        reverse=True,
    )[:30]
    ranked = rank_reference_candidates(candidates, result, search_terms)[:15]
    if len(ranked) < 10:
        raise RuntimeError(
            "Sumber ilmiah hanya menghasilkan "
            f"{len(ranked)} referensi dengan metadata lengkap. "
            "Minimal 10 referensi aktual diperlukan; sistem tidak membuat referensi fiktif."
        )
    return {
        "kata_kunci": search_terms["kata_kunci"],
        "query_pencarian": search_terms["query_pencarian"],
        "rekomendasi": ranked,
        "referensi_dipilih": [],
        "catatan_pencarian": (
            f"Ditemukan {len(ranked)} referensi aktual dari OpenAlex/Crossref. "
            "Status quartile Scopus/WoS atau SINTA perlu diverifikasi pada portal indeks terkait."
        ),
        "error_sumber": errors,
    }


def prepare_pdf_value(result_key, field_key, value):
    if result_key != "referensi_tambahan":
        return value
    if field_key == "rekomendasi":
        visible_keys = [
            "judul",
            "penulis",
            "tahun",
            "jurnal",
            "quartile_atau_sinta",
            "doi",
            "kategori",
            "alasan_relevan",
            "apa_7",
        ]
        return [
            {key: reference.get(key, "") for key in visible_keys}
            for reference in value
        ]
    if field_key == "referensi_dipilih":
        return [reference.get("apa_7", "") for reference in value]
    return value


def result_template_context(filename, output_name, result):
    return {
        "filename": filename,
        "output_name": output_name,
        "review_complete": is_review_complete(result),
        "workflow_steps": get_workflow_steps(result),
        "result": result,
    }


def get_workflow_steps(result=None):
    result = result or {}
    is_uploaded = bool(result.get("_uploaded"))
    unlocked_next = is_uploaded
    steps = []

    for step in WORKFLOW_STEPS:
        is_complete = bool(result.get(step["key"]))
        is_unlocked = is_complete or unlocked_next

        if not is_complete:
            unlocked_next = False
        else:
            unlocked_next = True

        if is_complete:
            status = "Selesai"
            icon = "✓"
        elif is_unlocked:
            status = "Aktif"
            icon = "▶"
        else:
            status = "Terkunci"
            icon = "🔒"

        steps.append(
            {
                **step,
                "complete": is_complete,
                "unlocked": is_unlocked,
                "status": status,
                "icon": icon,
            }
        )

    return steps


def module_is_unlocked(result, module_key):
    for step in get_workflow_steps(result):
        if step["key"] == module_key:
            return step["unlocked"]
    return False


def is_review_complete(result):
    return all(result.get(step["key"]) for step in WORKFLOW_STEPS)


def guard_module_order(result, module_key):
    if module_is_unlocked(result, module_key):
        return True

    flash("Tombol ini masih terkunci. Jalankan analisis sebelumnya sesuai urutan terlebih dahulu.")
    return False


def save_result(filename, result):
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    output_name = f"{Path(filename).stem}-{timestamp}.json"
    output_path = OUTPUT_DIR / output_name
    output_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return output_name


def load_result(output_name):
    output_path = OUTPUT_DIR / secure_filename(output_name)
    if not output_path.exists():
        return None
    return json.loads(output_path.read_text(encoding="utf-8"))


def update_result(output_name, result):
    output_path = OUTPUT_DIR / secure_filename(output_name)
    output_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")


def uploaded_file_is_used_elsewhere(filename, excluded_output_name):
    for output_path in OUTPUT_DIR.glob("*.json"):
        if output_path.name == excluded_output_name:
            continue
        try:
            other_result = json.loads(output_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if secure_filename(other_result.get("_filename", "")) == filename:
            return True
    return False


def pdf_safe_text(value):
    text = str(value)
    return escape(text).replace("\n", "<br/>")


def add_pdf_value(story, value, styles, level=0):
    if isinstance(value, dict):
        for key, child in value.items():
            label = str(key).replace("_", " ").title()
            story.append(Paragraph(pdf_safe_text(label), styles["FieldLabel"]))
            add_pdf_value(story, child, styles, level + 1)
        return

    if isinstance(value, list):
        if not value:
            story.append(Paragraph("Tidak ada.", styles["BodyText"]))
            return
        for item in value:
            if isinstance(item, (dict, list)):
                add_pdf_value(story, item, styles, level + 1)
            else:
                story.append(
                    Paragraph(
                        pdf_safe_text(item),
                        styles["ReviewBullet"],
                        bulletText="-",
                    )
                )
        return

    story.append(
        Paragraph(
            pdf_safe_text(value if value is not None else "Tidak ada."),
            styles["BodyText"],
        )
    )


def draw_pdf_page(canvas, document):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D9E2E8"))
    canvas.line(20 * mm, 15 * mm, 190 * mm, 15 * mm)
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#627085"))
    canvas.drawString(20 * mm, 10 * mm, "Laporan Hasil Reviu Artikel Ilmiah")
    canvas.drawRightString(190 * mm, 10 * mm, f"Halaman {document.page}")
    canvas.restoreState()


def build_review_pdf(result, pdf_path):
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ReportTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            textColor=colors.HexColor("#0F5147"),
            fontSize=20,
            leading=25,
            spaceAfter=12,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportMeta",
            parent=styles["Normal"],
            alignment=TA_CENTER,
            textColor=colors.HexColor("#627085"),
            fontSize=9,
            leading=13,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ModuleTitle",
            parent=styles["Heading1"],
            textColor=colors.HexColor("#146C5D"),
            fontSize=16,
            leading=20,
            spaceAfter=12,
        )
    )
    styles.add(
        ParagraphStyle(
            name="FieldTitle",
            parent=styles["Heading2"],
            textColor=colors.HexColor("#233142"),
            fontSize=11,
            leading=14,
            spaceBefore=9,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="FieldLabel",
            parent=styles["Normal"],
            textColor=colors.HexColor("#0F5147"),
            fontSize=9,
            leading=12,
            leftIndent=6 * mm,
            spaceBefore=5,
            spaceAfter=2,
            fontName="Helvetica-Bold",
        )
    )
    styles["BodyText"].fontName = "Helvetica"
    styles["BodyText"].fontSize = 9
    styles["BodyText"].leading = 13
    styles["BodyText"].spaceAfter = 6
    styles.add(
        ParagraphStyle(
            name="ReviewBullet",
            parent=styles["BodyText"],
            leftIndent=7 * mm,
            firstLineIndent=-4 * mm,
            bulletIndent=2 * mm,
        )
    )

    document = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=22 * mm,
        title="Laporan Hasil Reviu Artikel Ilmiah",
        author="Aplikasi Reviu Artikel Ilmiah AI",
    )
    filename = result.get("_filename", "Artikel ilmiah")
    created_at = result.get("_created_at", "")
    story = [
        Spacer(1, 12 * mm),
        Paragraph("Laporan Hasil Reviu Artikel Ilmiah", styles["ReportTitle"]),
        Paragraph(pdf_safe_text(filename), styles["ReportMeta"]),
        Paragraph(
            pdf_safe_text(f"Dibuat: {created_at}" if created_at else "Hasil akhir analisis"),
            styles["ReportMeta"],
        ),
        Spacer(1, 12 * mm),
    ]

    for section_index, (result_key, title, fields) in enumerate(PDF_REVIEW_SECTIONS):
        if section_index:
            story.append(Spacer(1, 6 * mm))
        story.append(Paragraph(f"{section_index + 1}. {pdf_safe_text(title)}", styles["ModuleTitle"]))
        module_data = result.get(result_key, {})
        analysis_data = module_data.get("hasil", module_data)

        for field_index, (field_key, field_title) in enumerate(fields):
            story.append(
                Paragraph(
                    f"{field_index + 1}. {pdf_safe_text(field_title)}",
                    styles["FieldTitle"],
                )
            )
            add_pdf_value(
                story,
                prepare_pdf_value(
                    result_key,
                    field_key,
                    analysis_data.get(field_key, "Belum ada hasil."),
                ),
                styles,
            )

    document.build(story, onFirstPage=draw_pdf_page, onLaterPages=draw_pdf_page)


@app.route("/", methods=["GET"])
def index():
    return render_template(
        "index.html",
        openai_warning=get_openai_warning(),
        workflow_steps=get_workflow_steps(),
    )


@app.route("/test-ai", methods=["POST"])
def test_ai():
    try:
        reply = test_openai_connection()
        flash(f"✅ OpenAI GPT-5.6 terhubung\nBalasan model: {reply}")
    except Exception as exc:
        flash(f"Tes koneksi AI gagal: {exc}")

    return redirect(url_for("index"))


@app.route("/download-review/<output_name>", methods=["GET"])
def download_review(output_name):
    safe_output_name = secure_filename(output_name)
    result = load_result(safe_output_name)
    if not result:
        flash("Hasil reviu tidak ditemukan.")
        return redirect(url_for("index"))

    if not is_review_complete(result):
        flash("PDF hanya dapat diunduh setelah semua tahap analisis selesai.")
        return redirect(url_for("index"))

    pdf_name = f"{Path(safe_output_name).stem}.pdf"
    pdf_path = OUTPUT_DIR / pdf_name
    build_review_pdf(result, pdf_path)
    return send_file(pdf_path, as_attachment=True, download_name=pdf_name)


@app.route("/delete-review/<output_name>", methods=["POST"])
def delete_review(output_name):
    safe_output_name = secure_filename(output_name)
    result = load_result(safe_output_name)
    if not result:
        flash("Hasil reviu tidak ditemukan atau sudah dihapus.")
        return redirect(url_for("index"))

    output_path = OUTPUT_DIR / safe_output_name
    pdf_path = OUTPUT_DIR / f"{Path(safe_output_name).stem}.pdf"
    uploaded_name = secure_filename(result.get("_filename", ""))

    output_path.unlink(missing_ok=True)
    pdf_path.unlink(missing_ok=True)
    if uploaded_name and not uploaded_file_is_used_elsewhere(uploaded_name, safe_output_name):
        uploaded_path = UPLOAD_DIR / uploaded_name
        uploaded_path.unlink(missing_ok=True)

    flash("Hasil reviu dan file terkait berhasil dihapus.")
    return redirect(url_for("index"))


@app.route("/analyze", methods=["POST"])
def analyze():
    uploaded_file = request.files.get("article")

    if not uploaded_file or uploaded_file.filename == "":
        flash("Silakan pilih file artikel .docx terlebih dahulu.")
        return redirect(url_for("index"))

    if not allowed_file(uploaded_file.filename):
        flash("Format file harus .docx.")
        return redirect(url_for("index"))

    filename = secure_filename(uploaded_file.filename)
    saved_path = UPLOAD_DIR / filename
    uploaded_file.save(saved_path)

    try:
        article_text = extract_docx_text(saved_path)
        if not article_text.strip():
            flash("Teks artikel tidak ditemukan di file .docx.")
            return redirect(url_for("index"))

        result = {
            "_uploaded": True,
            "_filename": filename,
            "_created_at": datetime.now().isoformat(),
        }
        output_name = save_result(filename, result)
    except Exception as exc:
        flash(str(exc))
        return redirect(url_for("index"))

    return render_template("result.html", **result_template_context(filename, output_name, result))


@app.route("/analyze-title", methods=["POST"])
def analyze_title():
    filename = secure_filename(request.form.get("filename", ""))
    output_name = secure_filename(request.form.get("output_name", ""))
    saved_path = UPLOAD_DIR / filename

    if not filename or not saved_path.exists():
        flash("File artikel tidak ditemukan. Silakan unggah artikel kembali.")
        return redirect(url_for("index"))

    result = load_result(output_name) if output_name else None
    if result is None:
        flash("Hasil analisis awal tidak ditemukan. Silakan analisis artikel dari awal.")
        return redirect(url_for("index"))
    if not guard_module_order(result, "judul"):
        return render_template("result.html", **result_template_context(filename, output_name, result))

    try:
        article_text = extract_docx_text(saved_path)
        title_info = detect_title(article_text)
        result["judul"] = analyze_title_with_ai(
            title_info["title"],
            build_short_article_context(article_text),
        )
        update_result(output_name, result)
    except Exception as exc:
        flash(str(exc))
        return render_template("result.html", **result_template_context(filename, output_name, result))

    return render_template("result.html", **result_template_context(filename, output_name, result))


@app.route("/analyze-abstract", methods=["POST"])
def analyze_abstract():
    filename = secure_filename(request.form.get("filename", ""))
    output_name = secure_filename(request.form.get("output_name", ""))
    saved_path = UPLOAD_DIR / filename

    if not filename or not saved_path.exists():
        flash("File artikel tidak ditemukan. Silakan unggah artikel kembali.")
        return redirect(url_for("index"))

    result = load_result(output_name) if output_name else None
    if result is None:
        flash("Hasil analisis awal tidak ditemukan. Silakan analisis artikel dari awal.")
        return redirect(url_for("index"))
    if not guard_module_order(result, "abstrak"):
        return render_template("result.html", **result_template_context(filename, output_name, result))

    try:
        article_text = extract_docx_text(saved_path)
        abstract_info = detect_abstract_section(article_text)
        result["abstrak"] = analyze_abstract_with_ai(
            abstract_info.get("text", ""),
            result["judul"],
            build_short_article_context(article_text),
        )
        update_result(output_name, result)
    except Exception as exc:
        flash(str(exc))
        return render_template("result.html", **result_template_context(filename, output_name, result))

    return render_template("result.html", **result_template_context(filename, output_name, result))


@app.route("/analyze-introduction", methods=["POST"])
def analyze_introduction():
    filename = secure_filename(request.form.get("filename", ""))
    output_name = secure_filename(request.form.get("output_name", ""))
    saved_path = UPLOAD_DIR / filename

    if not filename or not saved_path.exists():
        flash("File artikel tidak ditemukan. Silakan unggah artikel kembali.")
        return redirect(url_for("index"))

    result = load_result(output_name) if output_name else None
    if result is None:
        flash("Hasil analisis awal tidak ditemukan. Silakan analisis artikel dari awal.")
        return redirect(url_for("index"))
    if not guard_module_order(result, "pendahuluan"):
        return render_template("result.html", **result_template_context(filename, output_name, result))

    try:
        article_text = extract_docx_text(saved_path)
        introduction_info = detect_introduction_section(article_text)
        result["pendahuluan"] = analyze_introduction_with_ai(
            introduction_info.get("text", ""),
            result["judul"],
            result["abstrak"],
        )
        update_result(output_name, result)
    except Exception as exc:
        flash(str(exc))
        return render_template("result.html", **result_template_context(filename, output_name, result))

    return render_template("result.html", **result_template_context(filename, output_name, result))


@app.route("/analyze-methodology", methods=["POST"])
def analyze_methodology():
    filename = secure_filename(request.form.get("filename", ""))
    output_name = secure_filename(request.form.get("output_name", ""))
    saved_path = UPLOAD_DIR / filename

    if not filename or not saved_path.exists():
        flash("File artikel tidak ditemukan. Silakan unggah artikel kembali.")
        return redirect(url_for("index"))

    result = load_result(output_name) if output_name else None
    if result is None:
        flash("Hasil analisis awal tidak ditemukan. Silakan analisis artikel dari awal.")
        return redirect(url_for("index"))
    if not guard_module_order(result, "metodologi"):
        return render_template("result.html", **result_template_context(filename, output_name, result))

    try:
        article_text = extract_docx_text(saved_path)
        methodology_info = detect_methodology_section(article_text)
        result["metodologi"] = analyze_methodology_with_ai(
            methodology_info.get("text", ""),
            result["judul"],
            result["abstrak"],
            result["pendahuluan"],
        )
        update_result(output_name, result)
    except Exception as exc:
        flash(str(exc))
        return render_template("result.html", **result_template_context(filename, output_name, result))

    return render_template("result.html", **result_template_context(filename, output_name, result))


@app.route("/analyze-results-discussion", methods=["POST"])
def analyze_results_discussion():
    filename = secure_filename(request.form.get("filename", ""))
    output_name = secure_filename(request.form.get("output_name", ""))
    saved_path = UPLOAD_DIR / filename

    if not filename or not saved_path.exists():
        flash("File artikel tidak ditemukan. Silakan unggah artikel kembali.")
        return redirect(url_for("index"))

    result = load_result(output_name) if output_name else None
    if result is None:
        flash("Hasil analisis awal tidak ditemukan. Silakan analisis artikel dari awal.")
        return redirect(url_for("index"))
    if not guard_module_order(result, "hasil_pembahasan"):
        return render_template("result.html", **result_template_context(filename, output_name, result))

    try:
        article_text = extract_docx_text(saved_path)
        section_info = detect_results_discussion_section(article_text)
        result["hasil_pembahasan"] = analyze_results_discussion_with_ai(
            section_info.get("text", ""),
            result["judul"],
            result["abstrak"],
            result["pendahuluan"],
            result["metodologi"],
        )
        update_result(output_name, result)
    except Exception as exc:
        flash(str(exc))
        return render_template("result.html", **result_template_context(filename, output_name, result))

    return render_template("result.html", **result_template_context(filename, output_name, result))


@app.route("/analyze-conclusion", methods=["POST"])
def analyze_conclusion():
    filename = secure_filename(request.form.get("filename", ""))
    output_name = secure_filename(request.form.get("output_name", ""))
    saved_path = UPLOAD_DIR / filename

    if not filename or not saved_path.exists():
        flash("File artikel tidak ditemukan. Silakan unggah artikel kembali.")
        return redirect(url_for("index"))

    result = load_result(output_name) if output_name else None
    if result is None:
        flash("Hasil analisis awal tidak ditemukan. Silakan analisis artikel dari awal.")
        return redirect(url_for("index"))
    if not guard_module_order(result, "kesimpulan"):
        return render_template("result.html", **result_template_context(filename, output_name, result))

    try:
        article_text = extract_docx_text(saved_path)
        conclusion_info = detect_conclusion_section(article_text)
        result["kesimpulan"] = analyze_conclusion_with_ai(
            conclusion_info.get("text", ""),
            result,
        )
        update_result(output_name, result)
    except Exception as exc:
        flash(str(exc))

    return render_template("result.html", **result_template_context(filename, output_name, result))


def analyze_consistency():
    filename = secure_filename(request.form.get("filename", ""))
    output_name = secure_filename(request.form.get("output_name", ""))
    saved_path = UPLOAD_DIR / filename

    if not filename or not saved_path.exists():
        flash("File artikel tidak ditemukan. Silakan unggah artikel kembali.")
        return redirect(url_for("index"))

    result = load_result(output_name) if output_name else None
    if result is None:
        flash("Hasil analisis awal tidak ditemukan. Silakan analisis artikel dari awal.")
        return redirect(url_for("index"))
    if not guard_module_order(result, "konsistensi_artikel"):
        return render_template("result.html", **result_template_context(filename, output_name, result))

    try:
        article_text = extract_docx_text(saved_path)
        consistency_result = analyze_consistency_with_ai(article_text, result)
        result["konsistensi_artikel"] = {
            "hasil": consistency_result,
        }
        update_result(output_name, result)
    except Exception as exc:
        flash(str(exc))
        return render_template("result.html", **result_template_context(filename, output_name, result))

    return render_template("result.html", **result_template_context(filename, output_name, result))


def analyze_reviewer_simulator():
    filename = secure_filename(request.form.get("filename", ""))
    output_name = secure_filename(request.form.get("output_name", ""))
    saved_path = UPLOAD_DIR / filename

    if not filename or not saved_path.exists():
        flash("File artikel tidak ditemukan. Silakan unggah artikel kembali.")
        return redirect(url_for("index"))

    result = load_result(output_name) if output_name else None
    if result is None:
        flash("Hasil analisis awal tidak ditemukan. Silakan analisis artikel dari awal.")
        return redirect(url_for("index"))
    if not guard_module_order(result, "reviewer_jurnal_simulator"):
        return render_template("result.html", **result_template_context(filename, output_name, result))

    try:
        article_text = extract_docx_text(saved_path)
        reviewer_result = analyze_reviewer_simulator_with_ai(article_text, result)
        result["reviewer_jurnal_simulator"] = {
            "hasil": reviewer_result,
        }
        update_result(output_name, result)
    except Exception as exc:
        flash(str(exc))
        return render_template("result.html", **result_template_context(filename, output_name, result))

    return render_template("result.html", **result_template_context(filename, output_name, result))


@app.route("/analyze-reference-citation", methods=["POST"])
def analyze_reference_citation():
    filename = secure_filename(request.form.get("filename", ""))
    output_name = secure_filename(request.form.get("output_name", ""))
    saved_path = UPLOAD_DIR / filename

    if not filename or not saved_path.exists():
        flash("File artikel tidak ditemukan. Silakan unggah artikel kembali.")
        return redirect(url_for("index"))

    result = load_result(output_name) if output_name else None
    if result is None:
        flash("Hasil analisis awal tidak ditemukan. Silakan analisis artikel dari awal.")
        return redirect(url_for("index"))
    if not guard_module_order(result, "sitasi_apa"):
        return render_template("result.html", **result_template_context(filename, output_name, result))

    try:
        article_text = extract_docx_text(saved_path)
        reference_info = detect_reference_section(article_text)
        result["sitasi_apa"] = analyze_reference_citation_with_ai(
            article_text,
            reference_info.get("text", ""),
            result,
        )
        update_result(output_name, result)
    except Exception as exc:
        flash(str(exc))
        return render_template("result.html", **result_template_context(filename, output_name, result))

    return render_template("result.html", **result_template_context(filename, output_name, result))


@app.route("/analyze-reference-recommendations", methods=["POST"])
def analyze_reference_recommendations():
    filename = secure_filename(request.form.get("filename", ""))
    output_name = secure_filename(request.form.get("output_name", ""))
    saved_path = UPLOAD_DIR / filename

    if not filename or not saved_path.exists():
        flash("File artikel tidak ditemukan. Silakan unggah artikel kembali.")
        return redirect(url_for("index"))

    result = load_result(output_name) if output_name else None
    if result is None:
        flash("Hasil revisi tidak ditemukan. Silakan mulai kembali.")
        return redirect(url_for("index"))
    if not guard_module_order(result, "referensi_tambahan"):
        return render_template("result.html", **result_template_context(filename, output_name, result))

    try:
        article_text = extract_docx_text(saved_path)
        reference_info = detect_reference_section(article_text)
        recommendation_result = find_verified_reference_recommendations(result)
        recommendation_result["referensi_asli"] = reference_info.get("text", "")
        result["referensi_tambahan"] = recommendation_result
        update_result(output_name, result)
    except Exception as exc:
        flash(f"Pencarian referensi gagal: {exc}")

    return render_template("result.html", **result_template_context(filename, output_name, result))


@app.route("/add-recommended-reference", methods=["POST"])
def add_recommended_reference():
    filename = secure_filename(request.form.get("filename", ""))
    output_name = secure_filename(request.form.get("output_name", ""))
    doi = normalize_doi(request.form.get("doi", ""))
    result = load_result(output_name) if output_name else None

    if result is None or not result.get("referensi_tambahan"):
        flash("Hasil rekomendasi referensi tidak ditemukan.")
        return redirect(url_for("index"))

    recommendations = result["referensi_tambahan"].get("rekomendasi", [])
    selected_reference = next(
        (item for item in recommendations if normalize_doi(item.get("doi")) == doi),
        None,
    )
    if selected_reference is None:
        flash("Referensi tidak valid atau tidak berasal dari hasil pencarian.")
    else:
        selected = result["referensi_tambahan"].setdefault("referensi_dipilih", [])
        if not any(normalize_doi(item.get("doi")) == doi for item in selected):
            selected.append(selected_reference)
            update_result(output_name, result)
            flash("Referensi berhasil ditambahkan ke daftar pustaka pilihan.")
        else:
            flash("Referensi tersebut sudah ditambahkan.")

    return render_template(
        "result.html",
        **result_template_context(filename, output_name, result),
    )


def analyze_novelty_contribution():
    filename = secure_filename(request.form.get("filename", ""))
    output_name = secure_filename(request.form.get("output_name", ""))
    saved_path = UPLOAD_DIR / filename

    if not filename or not saved_path.exists():
        flash("File artikel tidak ditemukan. Silakan unggah artikel kembali.")
        return redirect(url_for("index"))

    result = load_result(output_name) if output_name else None
    if result is None:
        flash("Hasil analisis awal tidak ditemukan. Silakan analisis artikel dari awal.")
        return redirect(url_for("index"))
    if not guard_module_order(result, "novelty_dan_kontribusi"):
        return render_template("result.html", **result_template_context(filename, output_name, result))

    try:
        article_text = extract_docx_text(saved_path)
        novelty_result = analyze_novelty_contribution_with_ai(article_text, result)
        result["novelty_dan_kontribusi"] = {
            "hasil": novelty_result,
        }
        update_result(output_name, result)
    except Exception as exc:
        flash(str(exc))
        return render_template("result.html", **result_template_context(filename, output_name, result))

    return render_template("result.html", **result_template_context(filename, output_name, result))


def analyze_submission_readiness():
    filename = secure_filename(request.form.get("filename", ""))
    output_name = secure_filename(request.form.get("output_name", ""))
    saved_path = UPLOAD_DIR / filename

    if not filename or not saved_path.exists():
        flash("File artikel tidak ditemukan. Silakan unggah artikel kembali.")
        return redirect(url_for("index"))

    result = load_result(output_name) if output_name else None
    if result is None:
        flash("Hasil analisis awal tidak ditemukan. Silakan analisis artikel dari awal.")
        return redirect(url_for("index"))
    if not guard_module_order(result, "kesiapan_submit"):
        return render_template("result.html", **result_template_context(filename, output_name, result))

    try:
        article_text = extract_docx_text(saved_path)
        readiness_result = analyze_submission_readiness_with_ai(article_text, result)
        result["kesiapan_submit"] = {
            "hasil": readiness_result,
        }
        update_result(output_name, result)
    except Exception as exc:
        flash(str(exc))
        return render_template("result.html", **result_template_context(filename, output_name, result))

    return render_template("result.html", **result_template_context(filename, output_name, result))


if __name__ == "__main__":
    app.run(debug=True)
