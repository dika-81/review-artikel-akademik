"""Pembuat dokumen .docx sintetis untuk pengujian berbagai jenis artikel."""

import io

from docx import Document


def build_docx(paragraphs, tables=None, path=None):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    for table_rows in tables or []:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for col_index, cell in enumerate(row):
                table.cell(row_index, col_index).text = str(cell)
    if path is not None:
        document.save(str(path))
        return path
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


QUANTITATIVE_EXPERIMENT = [
    "Pengaruh Media Tangram terhadap Pemahaman Konsep Geometri Siswa Sekolah Dasar",
    "Abstrak",
    "Penelitian ini bertujuan menguji pengaruh media tangram terhadap pemahaman konsep geometri. "
    "Data dikumpulkan melalui angket. Hasil menunjukkan peningkatan yang signifikan.",
    "Kata kunci: tangram, geometri, sekolah dasar",
    "Pendahuluan",
    "Pemahaman konsep geometri siswa sekolah dasar masih rendah. "
    "Hasil observasi awal menunjukkan 60 persen siswa belum tuntas.",
    "Penelitian tentang media tangram belum banyak dilakukan. "
    "Tujuan penelitian ini adalah menguji pengaruh media tangram terhadap pemahaman konsep geometri.",
    "Metode Penelitian",
    "Penelitian ini menggunakan desain kuasi eksperimen dengan nonequivalent control group design. "
    "Sampel dipilih dengan purposive sampling, yaitu kelas III A sebagai kelas eksperimen dan kelas III B sebagai kelas kontrol.",
    "Instrumen penelitian terdiri atas tes pemahaman konsep, lembar observasi, dan angket respons. "
    "Data dianalisis menggunakan uji-t setelah uji normalitas dan uji homogenitas, serta dihitung N-gain dan effect size Cohen's d.",
    "Hasil dan Pembahasan",
    "Rata-rata posttest kelas eksperimen sebesar 82,4 sedangkan kelas kontrol 71,2. "
    "Hasil uji-t menunjukkan nilai p sebesar 0,001 sehingga terdapat perbedaan yang signifikan. "
    "Tabel 1 menyajikan hasil uji normalitas pretest.",
    "Temuan ini sejalan dengan penelitian Widodo (2021) tentang media manipulatif. "
    "Media tangram membantu siswa memvisualisasikan bangun datar (Sari & Putra, 2020).",
    "Kesimpulan",
    "Media tangram berpengaruh signifikan terhadap pemahaman konsep geometri siswa. "
    "Efektivitas ini didukung nilai N-gain kategori sedang.",
    "Daftar Pustaka",
    "Widodo, A. (2021). Media manipulatif dalam pembelajaran geometri. Jurnal Pendidikan Dasar, 12(1), 1-12.",
    "Sari, D., & Putra, R. (2020). Visualisasi bangun datar melalui tangram. Jurnal Matematika Sekolah, 8(2), 45-58.",
    "Rahman, T. (2019). Pembelajaran matematika realistik. Jurnal Elemen, 5(1), 20-31.",
]

QUALITATIVE_CASE_STUDY = [
    "Pengalaman Guru dalam Menerapkan Pembelajaran Diferensiasi: Studi Kasus",
    "Abstrak",
    "Penelitian kualitatif dengan pendekatan studi kasus ini bertujuan memahami pengalaman guru "
    "dalam menerapkan pembelajaran diferensiasi.",
    "Kata kunci: studi kasus, diferensiasi, guru",
    "Pendahuluan",
    "Pembelajaran diferensiasi menuntut pemahaman mendalam terhadap kebutuhan siswa. "
    "Penelitian ini mengeksplorasi pengalaman guru di satu sekolah dasar.",
    "Metode Penelitian",
    "Penelitian ini menggunakan pendekatan kualitatif dengan desain studi kasus. "
    "Data dikumpulkan melalui wawancara mendalam, observasi partisipan, dan dokumentasi terhadap lima informan guru.",
    "Data dianalisis secara tematik melalui pengodean, kategorisasi, dan perumusan tema. "
    "Keabsahan data dijaga melalui triangulasi sumber dan member checking.",
    "Hasil dan Pembahasan",
    "Tiga tema utama ditemukan: kesiapan guru, hambatan waktu, dan dukungan sekolah. "
    "Salah satu informan menyatakan bahwa waktu persiapan menjadi tantangan terbesar.",
    "Kesimpulan",
    "Pengalaman guru menunjukkan bahwa penerapan diferensiasi membutuhkan dukungan struktural sekolah.",
    "Daftar Pustaka",
    "Tomlinson, C. (2017). How to differentiate instruction. ASCD.",
]

SLR_ARTICLE = [
    "Tren Penelitian Media Pembelajaran Digital: Systematic Literature Review",
    "Abstrak",
    "Kajian ini merupakan systematic literature review terhadap 45 artikel tentang media pembelajaran digital.",
    "Kata kunci: systematic literature review, media digital",
    "Pendahuluan",
    "Perkembangan media digital menuntut pemetaan tren penelitian secara sistematis.",
    "Metode Penelitian",
    "Kajian ini mengikuti protokol PRISMA. Pencarian dilakukan pada basis data Scopus dan ERIC "
    "dengan string pencarian tertentu pada rentang 2015-2024.",
    "Kriteria inklusi dan eksklusi ditetapkan, dilanjutkan proses screening judul, abstrak, dan teks penuh, "
    "penghapusan duplikat, serta ekstraksi data.",
    "Hasil dan Pembahasan",
    "Dari 45 artikel, ditemukan tiga tren utama penggunaan media digital.",
    "Kesimpulan",
    "Tren penelitian media digital didominasi oleh studi pengembangan aplikasi mobile.",
    "Daftar Pustaka",
    "Page, M. (2021). The PRISMA 2020 statement. BMJ, 372, n71.",
]

NARRATIVE_REVIEW_ARTICLE = [
    "Pembelajaran Berbasis Permainan di Sekolah Dasar: Studi Literatur",
    "Abstrak",
    "Studi literatur ini menelaah manfaat dan keterbatasan pembelajaran berbasis permainan.",
    "Kata kunci: studi literatur, permainan, sekolah dasar",
    "Pendahuluan",
    "Pembelajaran berbasis permainan banyak digunakan untuk meningkatkan keterlibatan siswa.",
    "Metode Penelitian",
    "Kajian literatur dilakukan terhadap artikel terdahulu dan buku yang relevan. "
    "Sumber pustaka dipilih berdasarkan keterkaitannya dengan pembelajaran sekolah dasar "
    "dan disintesis melalui analisis deskriptif.",
    "Hasil dan Pembahasan",
    "Sintesis literatur menunjukkan pola peningkatan motivasi dan risiko distraksi.",
    "Kesimpulan",
    "Pembelajaran berbasis permainan bermanfaat bila tujuan dan mekanisme umpan balik dirancang jelas.",
    "Daftar Pustaka",
    "Plass, J., Homer, B., & Kinzer, C. (2015). Foundations of game-based learning. Educational Psychologist, 50(4), 258-283.",
]

CONCEPTUAL_ARTICLE = [
    "Rekonseptualisasi Literasi Numerasi dalam Kurikulum Sekolah Dasar: Kajian Teoretis",
    "Abstrak",
    "Artikel konseptual ini menawarkan kerangka teoretis literasi numerasi berbasis konteks budaya.",
    "Kata kunci: literasi numerasi, kajian teoretis",
    "Pendahuluan",
    "Konsep literasi numerasi sering dipahami secara sempit sebagai keterampilan berhitung.",
    "Pembahasan",
    "Artikel ini merumuskan tiga proposisi tentang hubungan numerasi dan konteks budaya. "
    "Kerangka argumen dibangun dari telaah konsep pada literatur terdahulu.",
    "Kesimpulan",
    "Kerangka konseptual yang diusulkan memperluas batas penerapan konsep literasi numerasi.",
    "Daftar Pustaka",
    "Geiger, V. (2015). A rich interpretation of numeracy. ZDM, 47(4), 531-544.",
]

MIXED_METHODS_ARTICLE = [
    "Efektivitas dan Persepsi Siswa terhadap Pembelajaran Berbasis Proyek: Studi Mixed Methods",
    "Abstrak",
    "Penelitian mixed methods dengan desain explanatory sequential ini menggabungkan data kuantitatif dan kualitatif.",
    "Pendahuluan",
    "Pembelajaran berbasis proyek diduga meningkatkan hasil belajar sekaligus keterlibatan siswa.",
    "Metode Penelitian",
    "Desain explanatory sequential digunakan: tahap pertama survei dengan kuesioner kepada 120 siswa "
    "yang dianalisis dengan regresi, tahap kedua wawancara mendalam terhadap 8 informan.",
    "Integrasi data dilakukan pada tahap interpretasi melalui joint display.",
    "Hasil dan Pembahasan",
    "Hasil kuantitatif menunjukkan pengaruh signifikan; data wawancara memperdalam temuan tersebut.",
    "Kesimpulan",
    "Integrasi kedua jenis data menunjukkan efektivitas pembelajaran berbasis proyek.",
    "Daftar Pustaka",
    "Creswell, J., & Plano Clark, V. (2018). Designing and conducting mixed methods research. Sage.",
]

PTK_ARTICLE = [
    "Peningkatan Aktivitas Belajar melalui Model Discovery Learning: Penelitian Tindakan Kelas",
    "Abstrak",
    "Penelitian tindakan kelas ini dilaksanakan dalam dua siklus untuk meningkatkan aktivitas belajar siswa.",
    "Pendahuluan",
    "Aktivitas belajar siswa kelas IV masih rendah berdasarkan observasi awal.",
    "Metode Penelitian",
    "Penelitian tindakan kelas model Kemmis dan McTaggart dengan tahapan perencanaan, tindakan, observasi, "
    "dan refleksi pada setiap siklus. Indikator keberhasilan ditetapkan sebesar 80 persen.",
    "Hasil dan Pembahasan",
    "Pada siklus I ketuntasan mencapai 65 persen dan meningkat menjadi 86 persen pada siklus II.",
    "Kesimpulan",
    "Model discovery learning meningkatkan aktivitas belajar siswa hingga melampaui indikator keberhasilan.",
    "Daftar Pustaka",
    "Kemmis, S., & McTaggart, R. (2014). The action research planner. Springer.",
]

RND_ARTICLE = [
    "Pengembangan Media Interaktif Berbasis ADDIE untuk Pembelajaran IPA",
    "Abstrak",
    "Penelitian dan pengembangan ini menghasilkan media interaktif menggunakan model ADDIE.",
    "Pendahuluan",
    "Media pembelajaran IPA yang tersedia belum memenuhi kebutuhan siswa.",
    "Metode Penelitian",
    "Penelitian ini merupakan research and development dengan model ADDIE: analysis, design, development, "
    "implementation, dan evaluation. Produk divalidasi melalui validasi ahli materi dan ahli media, "
    "kemudian diuji coba terbatas menggunakan angket respon serta lembar validasi untuk menilai kelayakan dan kepraktisan.",
    "Hasil dan Pembahasan",
    "Skor validasi ahli materi sebesar 3,6 dari 4 dengan kategori sangat layak.",
    "Kesimpulan",
    "Media interaktif yang dikembangkan dinyatakan layak dan praktis digunakan.",
    "Daftar Pustaka",
    "Branch, R. (2009). Instructional design: The ADDIE approach. Springer.",
]

PHENOMENOLOGY_ARTICLE = [
    "Makna Pengalaman Belajar Daring bagi Mahasiswa Tahun Pertama: Studi Fenomenologi",
    "Abstrak",
    "Studi fenomenologi ini mengungkap makna pengalaman belajar daring mahasiswa.",
    "Pendahuluan",
    "Pandemi mengubah pengalaman belajar mahasiswa secara mendasar.",
    "Metode Penelitian",
    "Pendekatan fenomenologi digunakan dengan wawancara mendalam terhadap sepuluh partisipan. "
    "Analisis dilakukan melalui bracketing, horizonalisasi, dan perumusan esensi pengalaman.",
    "Hasil dan Pembahasan",
    "Esensi pengalaman partisipan terangkum dalam tema keterasingan dan adaptasi.",
    "Kesimpulan",
    "Makna pengalaman belajar daring berpusat pada proses adaptasi sosial-akademik.",
    "Daftar Pustaka",
    "Moustakas, C. (1994). Phenomenological research methods. Sage.",
]

SURVEY_ARTICLE = [
    "Hubungan antara Efikasi Diri dan Prokrastinasi Akademik Mahasiswa",
    "Abstrak",
    "Penelitian korelasional ini menguji hubungan antara efikasi diri dan prokrastinasi akademik.",
    "Pendahuluan",
    "Prokrastinasi akademik berkaitan dengan berbagai faktor psikologis.",
    "Metode Penelitian",
    "Penelitian survei korelasional dengan kuesioner skala likert kepada 210 mahasiswa. "
    "Data dianalisis dengan korelasi pearson dan regresi sederhana.",
    "Hasil dan Pembahasan",
    "Terdapat hubungan negatif signifikan antara efikasi diri dan prokrastinasi (r = -0,42; p < 0,05).",
    "Kesimpulan",
    "Efikasi diri berkontribusi menurunkan prokrastinasi akademik mahasiswa.",
    "Daftar Pustaka",
    "Bandura, A. (1997). Self-efficacy: The exercise of control. Freeman.",
]

THEMATIC_INTERVIEW_ARTICLE = [
    "Analisis Tematik Persepsi Guru terhadap Kurikulum Merdeka",
    "Abstrak",
    "Penelitian kualitatif ini menganalisis persepsi guru melalui analisis tematik berbasis wawancara.",
    "Pendahuluan",
    "Implementasi Kurikulum Merdeka menimbulkan beragam persepsi guru.",
    "Metode Penelitian",
    "Data dikumpulkan melalui wawancara semi terstruktur terhadap dua belas informan guru. "
    "Analisis tematik dilakukan melalui koding awal, pengelompokan kategori, dan perumusan tema. "
    "Triangulasi dan member checking dilakukan untuk menjaga kredibilitas.",
    "Hasil dan Pembahasan",
    "Empat tema persepsi guru berhasil diidentifikasi.",
    "Kesimpulan",
    "Persepsi guru bervariasi menurut dukungan sekolah dan pengalaman pelatihan.",
    "Daftar Pustaka",
    "Braun, V., & Clarke, V. (2006). Using thematic analysis in psychology. Qualitative Research in Psychology, 3(2), 77-101.",
]

ENGLISH_ARTICLE = [
    "The Effect of Game-Based Learning on Students' Motivation",
    "Abstract",
    "This quasi-experimental study examines the effect of game-based learning on motivation.",
    "Keywords: game-based learning, motivation",
    "Introduction",
    "Student motivation remains a persistent challenge in primary classrooms.",
    "Research Method",
    "A quasi-experimental design with an experimental group and a control group was used. "
    "Data were collected through a motivation questionnaire and analyzed using a t-test after normality testing.",
    "Results and Discussion",
    "The experimental group outperformed the control group significantly (p = 0.003).",
    "Conclusion",
    "Game-based learning significantly improves students' motivation.",
    "References",
    "Deci, E., & Ryan, R. (2000). The what and why of goal pursuits. Psychological Inquiry, 11(4), 227-268.",
]

NO_HEADING_ARTICLE = [
    "Catatan Reflektif Pembelajaran Kelas Rangkap",
    "Tulisan ini membahas praktik pembelajaran kelas rangkap di sekolah kecil. "
    "Pengalaman guru menjadi sumber utama refleksi.",
    "Refleksi menunjukkan pentingnya fleksibilitas kurikulum bagi kelas rangkap.",
]
