import io
import inspect as py_inspect
import json
import logging
import os
from pathlib import Path
import re
import subprocess
import sys

import httpx
import openai
import pytest
from docx import Document
from flask import Flask, render_template, session as flask_session
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import inspect

import app as app_module
from conftest import register_and_login, set_valid_api_key

FAKE_KEY = "sk-test-kunci-palsu-123"


def extract_csrf_token(html):
    match = re.search(
        rb'name="csrf_token"\s+value="([^"]+)"',
        html,
    )
    assert match is not None, "csrf_token tidak ditemukan di HTML"
    return match.group(1).decode()


def make_docx_bytes():
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Judul Artikel Uji Coba Media Pembelajaran")
    document.add_paragraph("Abstract")
    document.add_paragraph("Ini adalah abstrak penelitian uji coba.")
    document.add_paragraph("Introduction")
    document.add_paragraph("Ini pendahuluan.")
    document.save(buffer)
    buffer.seek(0)
    return buffer


def make_openai_error(cls, status, body):
    request = httpx.Request("POST", "https://api.openai.com/v1/responses")
    response = httpx.Response(status, request=request, json=body)
    return cls("boom", response=response, body=body.get("error"))


# 1. Belum login tidak dapat membuka halaman analisis.
def test_index_requires_login(client):
    response = client.get("/")
    assert response.status_code == 302
    assert "/login" in response.headers["Location"]


def test_analyze_requires_login(client):
    response = client.post("/analyze", data={})
    assert response.status_code == 302
    assert "/login" in response.headers["Location"]


# 2. Login tanpa API key diarahkan ke halaman API key.
def test_login_without_api_key_redirects_to_api_key_page(client):
    register_and_login(client)
    response = client.get("/")
    assert response.status_code == 302
    assert "/api-key" in response.headers["Location"]


def test_registration_commits_user_to_database(client):
    response = client.post(
        "/register",
        data={
            "username": "persisten",
            "password": "rahasia-123",
            "confirm_password": "rahasia-123",
        },
    )
    assert response.status_code == 302
    with app_module.app.app_context():
        app_module.db.session.remove()
        user = app_module.find_user_row("persisten")
        assert user is not None
        assert user.password_hash


def test_init_db_is_idempotent_and_keeps_existing_data(
    flask_app, client, monkeypatch
):
    client.post(
        "/register",
        data={
            "username": "tetap-ada",
            "password": "rahasia-123",
            "confirm_password": "rahasia-123",
        },
    )
    client.post(
        "/login",
        data={"username": "tetap-ada", "password": "rahasia-123"},
    )
    set_valid_api_key(client, monkeypatch, "sk-tetap-ada-3333")
    runner = flask_app.test_cli_runner()
    assert runner.invoke(args=["init-db"]).exit_code == 0
    assert runner.invoke(args=["init-db"]).exit_code == 0
    with flask_app.app_context():
        app_module.db.session.remove()
        user = app_module.find_user_row("tetap-ada")
        assert user is not None
        assert app_module.find_user_api_key_record(user.id) is not None


# 3. API key invalid ditolak.
def test_invalid_api_key_rejected(client, monkeypatch):
    register_and_login(client)
    monkeypatch.setattr(
        app_module,
        "validate_api_key",
        lambda api_key: "API key OpenAI tidak valid atau sudah dicabut. Periksa kembali API key Anda.",
    )
    response = client.post("/api-key", data={"api_key": "sk-salah"})
    assert "tidak valid".encode() in response.data
    with app_module.app.app_context():
        user = app_module.find_user_row("tester")
        assert app_module.find_user_api_key_record(user.id) is None


# 4 & 5. API key valid disimpan terenkripsi dan tidak bocor ke client/session.
def test_valid_api_key_stored_as_ciphertext_only(client, monkeypatch):
    register_and_login(client)
    response = set_valid_api_key(client, monkeypatch, FAKE_KEY)
    assert response.status_code == 302

    follow = client.get("/")
    assert FAKE_KEY.encode() not in follow.data
    for cookie in client._cookies.values():
        assert FAKE_KEY not in cookie.value
    with client.session_transaction() as flask_session:
        assert FAKE_KEY not in repr(dict(flask_session))

    api_page = client.get("/api-key")
    assert FAKE_KEY.encode() not in api_page.data
    assert FAKE_KEY[-4:].encode() in api_page.data

    with app_module.app.app_context():
        user = app_module.find_user_row("tester")
        record = app_module.find_user_api_key_record(user.id)
        assert record is not None
        assert isinstance(record.ciphertext, bytes)
        assert FAKE_KEY.encode() not in record.ciphertext
        assert record.last_four == FAKE_KEY[-4:]
        assert (
            app_module.API_KEY_CIPHER.decrypt(record.ciphertext).decode()
            == FAKE_KEY
        )


def test_administrator_openai_key_is_ignored(client, monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "sk-admin-yang-tidak-boleh-dipakai")
    register_and_login(client)
    response = client.get("/")
    assert response.status_code == 302
    assert "/api-key" in response.headers["Location"]
    client.get("/api-key")
    assert app_module.get_current_user_api_key() is None


# 6. Analisis memakai API key pengguna yang sedang aktif.
def test_analysis_uses_user_api_key(client, monkeypatch):
    captured = {}

    class FakeResponses:
        def create(self, **kwargs):
            class R:
                output_text = "Halo dari model"

            return R()

    class FakeClient:
        def __init__(self):
            self.responses = FakeResponses()

    def fake_build_user_client(api_key):
        captured["api_key"] = api_key
        return FakeClient()

    register_and_login(client)
    set_valid_api_key(client, monkeypatch, FAKE_KEY)
    monkeypatch.setattr(app_module, "build_user_client", fake_build_user_client)

    response = client.post("/test-ai", follow_redirects=True)
    assert response.status_code == 200
    assert captured["api_key"] == FAKE_KEY


# 7. Dua pengguna tidak berbagi session, key, maupun hasil.
def test_two_users_are_isolated(flask_app, monkeypatch):
    client_a = flask_app.test_client()
    client_b = flask_app.test_client()

    register_and_login(client_a, username="usera")
    register_and_login(client_b, username="userb")
    key_a = "sk-milik-user-a-AAAA"
    key_b = "sk-milik-user-b-BBBB"
    set_valid_api_key(client_a, monkeypatch, key_a)
    set_valid_api_key(client_b, monkeypatch, key_b)

    with app_module.app.app_context():
        user_a = app_module.find_user_row("usera")
        user_b = app_module.find_user_row("userb")
        record_a = app_module.find_user_api_key_record(user_a.id)
        record_b = app_module.find_user_api_key_record(user_b.id)
        assert record_a.user_id != record_b.user_id
        assert record_a.ciphertext != record_b.ciphertext
        assert app_module.API_KEY_CIPHER.decrypt(record_a.ciphertext).decode() == key_a
        assert app_module.API_KEY_CIPHER.decrypt(record_b.ciphertext).decode() == key_b

    page_a = client_a.get("/api-key")
    page_b = client_b.get("/api-key")
    assert b"AAAA" in page_a.data and b"BBBB" not in page_a.data
    assert b"BBBB" in page_b.data and b"AAAA" not in page_b.data

    # User A mengunggah artikel; user B tidak dapat memakai hasilnya.
    upload = client_a.post(
        "/analyze",
        data={"article": (make_docx_bytes(), "artikel.docx")},
        content_type="multipart/form-data",
    )
    assert upload.status_code == 200
    output_name = next(
        path.name
        for path in app_module.OUTPUT_DIR.glob("*/artikel-*.json")
    )
    response = client_b.post(
        "/analyze-title",
        data={"filename": "artikel.docx", "output_name": output_name},
        follow_redirects=True,
    )
    assert "tidak ditemukan".encode() in response.data


def test_each_openai_request_uses_the_logged_in_users_key(flask_app, monkeypatch):
    captured_keys = []

    class FakeResponses:
        def create(self, **kwargs):
            class Response:
                output_text = "ok"

            return Response()

    class FakeClient:
        responses = FakeResponses()

    def fake_build_user_client(api_key):
        captured_keys.append(api_key)
        return FakeClient()

    monkeypatch.setattr(app_module, "build_user_client", fake_build_user_client)
    client_a = flask_app.test_client()
    client_b = flask_app.test_client()
    register_and_login(client_a, username="request-a")
    register_and_login(client_b, username="request-b")
    set_valid_api_key(client_a, monkeypatch, "sk-request-user-a-1111")
    set_valid_api_key(client_b, monkeypatch, "sk-request-user-b-2222")

    client_a.post("/test-ai")
    client_b.post("/test-ai")
    assert captured_keys == [
        "sk-request-user-a-1111",
        "sk-request-user-b-2222",
    ]


def test_same_plaintext_produces_isolated_random_ciphertexts(flask_app, monkeypatch):
    client_a = flask_app.test_client()
    client_b = flask_app.test_client()
    register_and_login(client_a, username="cipher-a")
    register_and_login(client_b, username="cipher-b")
    same_key = "sk-sama-untuk-uji-4444"
    set_valid_api_key(client_a, monkeypatch, same_key)
    set_valid_api_key(client_b, monkeypatch, same_key)

    with app_module.app.app_context():
        user_a = app_module.find_user_row("cipher-a")
        user_b = app_module.find_user_row("cipher-b")
        record_a = app_module.find_user_api_key_record(user_a.id)
        record_b = app_module.find_user_api_key_record(user_b.id)
        assert record_a.user_id != record_b.user_id
        assert record_a.ciphertext != record_b.ciphertext
        assert same_key.encode() not in record_a.ciphertext
        assert same_key.encode() not in record_b.ciphertext


# 8 & 9. Logout mempertahankan key terenkripsi; login ulang langsung siap.
def test_logout_and_relogin_preserve_encrypted_api_key(client, monkeypatch):
    register_and_login(client)
    set_valid_api_key(client, monkeypatch, FAKE_KEY)

    response = client.post("/logout")
    assert "/login" in response.headers["Location"]
    with app_module.app.app_context():
        user = app_module.find_user_row("tester")
        assert user is not None
        assert app_module.find_user_api_key_record(user.id) is not None
    with client.session_transaction() as flask_session:
        assert "_user_id" not in flask_session
        assert FAKE_KEY not in repr(dict(flask_session))

    # Setelah logout, halaman analisis tidak dapat diakses lagi (tombol Back).
    assert "/login" in client.get("/").headers["Location"]

    login_response = client.post(
        "/login", data={"username": "tester", "password": "rahasia-123"}
    )
    assert login_response.headers["Location"].endswith("/")
    response = client.get("/")
    assert response.status_code == 200


# 10. Menghapus key tidak menghapus akun dan analisis kembali meminta key.
def test_user_can_replace_and_delete_own_api_key(client, monkeypatch):
    register_and_login(client)
    set_valid_api_key(client, monkeypatch, FAKE_KEY)
    with app_module.app.app_context():
        user = app_module.find_user_row("tester")
        first_ciphertext = app_module.find_user_api_key_record(user.id).ciphertext

    replacement_key = "sk-pengganti-sendiri-9876"
    set_valid_api_key(client, monkeypatch, replacement_key)
    with app_module.app.app_context():
        user = app_module.find_user_row("tester")
        replacement = app_module.find_user_api_key_record(user.id)
        assert replacement.ciphertext != first_ciphertext
        assert replacement.last_four == "9876"
        assert (
            app_module.API_KEY_CIPHER.decrypt(replacement.ciphertext).decode()
            == replacement_key
        )

    delete_response = client.post("/api-key/delete")
    assert delete_response.status_code == 302
    with app_module.app.app_context():
        user = app_module.find_user_row("tester")
        assert user is not None
        assert app_module.find_user_api_key_record(user.id) is None

    response = client.get("/")
    assert "/api-key" in response.headers["Location"]


def test_user_cannot_delete_another_users_api_key(flask_app, monkeypatch):
    client_a = flask_app.test_client()
    client_b = flask_app.test_client()
    register_and_login(client_a, username="pemilik-a")
    register_and_login(client_b, username="pemilik-b")
    set_valid_api_key(client_a, monkeypatch, "sk-user-a-1111")
    set_valid_api_key(client_b, monkeypatch, "sk-user-b-2222")

    client_a.post("/api-key/delete")

    with app_module.app.app_context():
        user_a = app_module.find_user_row("pemilik-a")
        user_b = app_module.find_user_row("pemilik-b")
        assert app_module.find_user_api_key_record(user_a.id) is None
        assert app_module.find_user_api_key_record(user_b.id) is not None


def test_api_key_never_appears_in_application_logs(client, monkeypatch, caplog):
    secret_value = "sk-jangan-masuk-log-LOG9"
    caplog.set_level("INFO")
    register_and_login(client)
    set_valid_api_key(client, monkeypatch, secret_value)
    client.get("/api-key")
    client.post("/logout")
    assert secret_value not in caplog.text
    for logger_name in ("openai", "httpx", "httpcore"):
        assert logging.getLogger(logger_name).level >= logging.WARNING


# 11. Terjemahan error OpenAI.
def test_translate_insufficient_quota():
    exc = make_openai_error(
        openai.RateLimitError,
        429,
        {"error": {"message": "quota", "type": "insufficient_quota", "code": "insufficient_quota"}},
    )
    message = app_module.translate_openai_error(exc)
    assert "Kuota atau saldo OpenAI API Anda tidak mencukupi" in message


def test_translate_rate_limit_not_quota():
    exc = make_openai_error(
        openai.RateLimitError,
        429,
        {"error": {"message": "slow down", "type": "rate_limit_error", "code": "rate_limit_exceeded"}},
    )
    message = app_module.translate_openai_error(exc)
    assert "rate limit" in message.lower()
    assert "saldo" not in message.lower()


def test_translate_invalid_key():
    exc = make_openai_error(
        openai.AuthenticationError,
        401,
        {"error": {"message": "bad key", "type": "invalid_request_error", "code": "invalid_api_key"}},
    )
    assert "tidak valid" in app_module.translate_openai_error(exc)


def test_translate_model_not_found():
    exc = make_openai_error(
        openai.NotFoundError,
        404,
        {"error": {"message": "no model", "type": "invalid_request_error", "code": "model_not_found"}},
    )
    assert "tidak tersedia" in app_module.translate_openai_error(exc)


def test_translate_timeout_and_connection():
    request = httpx.Request("POST", "https://api.openai.com/v1/responses")
    timeout_exc = openai.APITimeoutError(request=request)
    assert "batas waktu" in app_module.translate_openai_error(timeout_exc)
    connection_exc = openai.APIConnectionError(message="down", request=request)
    assert "Koneksi ke OpenAI gagal" in app_module.translate_openai_error(connection_exc)


def test_safe_error_message_hides_internal_details():
    assert "internal" in app_module.safe_error_message(ValueError("secret traceback info"))
    assert "secret" not in app_module.safe_error_message(ValueError("secret traceback info"))


# 12. File sementara dihapus setelah hasil reviu dihapus.
def test_delete_review_removes_files(client, monkeypatch):
    register_and_login(client)
    set_valid_api_key(client, monkeypatch, FAKE_KEY)
    client.post(
        "/analyze",
        data={"article": (make_docx_bytes(), "artikel.docx")},
        content_type="multipart/form-data",
    )
    output_path = next(app_module.OUTPUT_DIR.glob("*/artikel-*.json"))
    upload_path = next(app_module.UPLOAD_DIR.glob("*/artikel.docx"))

    client.post(f"/delete-review/{output_path.name}")
    assert not output_path.exists()
    assert not upload_path.exists()


# 13. Fitur lama: upload artikel valid tetap bekerja dan format salah ditolak.
def test_upload_docx_still_works(client, monkeypatch):
    register_and_login(client)
    set_valid_api_key(client, monkeypatch, FAKE_KEY)
    response = client.post(
        "/analyze",
        data={"article": (make_docx_bytes(), "artikel.docx")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    assert "Artikel berhasil diunggah".encode() in response.data


def test_upload_wrong_extension_rejected(client, monkeypatch):
    register_and_login(client)
    set_valid_api_key(client, monkeypatch, FAKE_KEY)
    response = client.post(
        "/analyze",
        data={"article": (io.BytesIO(b"bukan docx"), "artikel.exe")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert "Format file harus .docx".encode() in response.data


def test_upload_corrupt_docx_rejected(client, monkeypatch):
    register_and_login(client)
    set_valid_api_key(client, monkeypatch, FAKE_KEY)
    response = client.post(
        "/analyze",
        data={"article": (io.BytesIO(b"PK\x03\x04 rusak"), "artikel.docx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert "tidak dapat dibaca".encode() in response.data


# CSRF aktif menolak POST tanpa token.
def test_csrf_protects_post(flask_app):
    flask_app.config.update(WTF_CSRF_ENABLED=True)
    try:
        raw_client = flask_app.test_client()
        response = raw_client.post(
            "/login", data={"username": "x", "password": "y"}
        )
        assert response.status_code == 400
        assert "Sesi form sudah tidak berlaku".encode() in response.data
    finally:
        flask_app.config.update(WTF_CSRF_ENABLED=False)


def test_login_post_with_csrf_is_not_rejected(flask_app):
    flask_app.config.update(WTF_CSRF_ENABLED=True)
    try:
        raw_client = flask_app.test_client()
        token = extract_csrf_token(raw_client.get("/login").data)
        response = raw_client.post(
            "/login",
            data={"username": "tidak-ada", "password": "salah", "csrf_token": token},
        )
        assert response.status_code == 200
        assert "CSRF".encode() not in response.data
    finally:
        flask_app.config.update(WTF_CSRF_ENABLED=False)


def test_all_post_forms_include_csrf_token():
    template_dir = Path(app_module.BASE_DIR) / "templates"
    for template_path in template_dir.glob("*.html"):
        html = template_path.read_text(encoding="utf-8")
        for form_match in re.finditer(
            r"<form\b(?P<attrs>[^>]*)>(?P<body>.*?)</form>",
            html,
            flags=re.IGNORECASE | re.DOTALL,
        ):
            attrs = form_match.group("attrs")
            if re.search(r'method\s*=\s*["\']post["\']', attrs, re.IGNORECASE):
                assert "csrf_token" in form_match.group("body"), (
                    f"Form POST tanpa CSRF di {template_path.name}"
                )


def test_analyze_abstract_form_contains_csrf_token(flask_app):
    result = {
        "_uploaded": True,
        "_filename": "artikel.docx",
        "_created_at": "2026-07-19T00:00:00+00:00",
        "_profil_heuristik": {"jenis_artikel": "kuantitatif"},
        "_jenis_artikel_pilihan": "otomatis",
        "_mode_analisis": "ringkas",
        "profil_artikel": {"klasifikasi_artikel": {"jenis_artikel": "kuantitatif"}},
        "judul": {"judul_awal": "Judul uji", "judul_rekomendasi": "Judul uji"},
    }
    with flask_app.test_request_context("/hasil"):
        html = render_template(
            "result.html",
            **app_module.result_template_context(
                "artikel.docx", "artikel.json", result
            ),
        )
    abstract_form = re.search(
        r'<form[^>]+action="/analyze-abstract"[^>]*>(.*?)</form>',
        html,
        flags=re.DOTALL,
    )
    assert abstract_form is not None
    assert 'name="csrf_token"' in abstract_form.group(1)


def test_analyze_abstract_post_with_csrf_reaches_route(flask_app, monkeypatch):
    raw_client = flask_app.test_client()
    register_and_login(raw_client)
    set_valid_api_key(raw_client, monkeypatch, FAKE_KEY)

    flask_app.config.update(WTF_CSRF_ENABLED=True)
    try:
        token = extract_csrf_token(raw_client.get("/").data)
        response = raw_client.post(
            "/analyze-abstract",
            data={
                "csrf_token": token,
                "filename": "",
                "output_name": "",
            },
        )
        assert response.status_code == 302
        assert response.headers["Location"].endswith("/")
    finally:
        flask_app.config.update(WTF_CSRF_ENABLED=False)


def test_analyze_abstract_get_redirects_to_form_page(client, monkeypatch):
    register_and_login(client)
    set_valid_api_key(client, monkeypatch, FAKE_KEY)
    response = client.get("/analyze-abstract")
    assert response.status_code == 302
    assert response.headers["Location"].endswith("/")


def test_postgres_url_is_normalized():
    assert (
        app_module.normalize_database_url("postgres://user:pass@db/name")
        == "postgresql://user:pass@db/name"
    )
    assert (
        app_module.get_sqlalchemy_database_url("postgres://user:pass@db/name")
        == "postgresql+psycopg://user:pass@db/name"
    )


def test_production_never_falls_back_to_sqlite(monkeypatch):
    monkeypatch.setattr(app_module, "IS_PRODUCTION", True)
    monkeypatch.delenv("DATABASE_URL", raising=False)
    with pytest.raises(RuntimeError, match="DATABASE_URL PostgreSQL wajib"):
        app_module.get_database_url()

    monkeypatch.setenv("DATABASE_URL", "sqlite:///local.db")
    with pytest.raises(RuntimeError, match="bukan SQLite"):
        app_module.get_database_url()


def test_production_session_never_uses_ephemeral_filesystem(monkeypatch):
    monkeypatch.setattr(app_module, "IS_PRODUCTION", True)
    monkeypatch.delenv("SESSION_TYPE", raising=False)
    assert app_module.get_session_type() == "sqlalchemy"

    monkeypatch.setenv("SESSION_TYPE", "filesystem")
    with pytest.raises(RuntimeError, match="filesystem ephemeral"):
        app_module.get_session_type()


def test_api_key_encryption_key_is_required_and_valid_in_production(monkeypatch):
    monkeypatch.setattr(app_module, "IS_PRODUCTION", True)
    monkeypatch.delenv("API_KEY_ENCRYPTION_KEY", raising=False)
    with pytest.raises(RuntimeError, match="API_KEY_ENCRYPTION_KEY"):
        app_module.get_api_key_encryption_key("secret")

    monkeypatch.setenv("API_KEY_ENCRYPTION_KEY", "bukan-fernet-key")
    with pytest.raises(RuntimeError, match="tidak valid"):
        app_module.get_api_key_encryption_key("secret")


def test_deployment_config_is_byok_database_agnostic_and_initializes_schema():
    app_source = (Path(app_module.BASE_DIR) / "app.py").read_text(
        encoding="utf-8"
    )
    render_config = (Path(app_module.BASE_DIR) / "render.yaml").read_text(
        encoding="utf-8"
    )
    assert "OPENAI_API_KEY" not in app_source
    assert "API_KEY_SESSION_FIELD" not in app_source
    assert "OPENAI_API_KEY" not in render_config
    assert "API_KEY_ENCRYPTION_KEY" in render_config
    assert "fromDatabase" not in render_config
    assert "\ndatabases:" not in render_config
    build_command = next(
        line.strip()
        for line in render_config.splitlines()
        if line.strip().startswith("buildCommand:")
    )
    assert build_command == (
        "buildCommand: pip install -r requirements.txt "
        "&& flask --app app init-db"
    )
    lowered_build_command = build_command.lower()
    assert "printenv" not in lowered_build_command
    assert "env |" not in lowered_build_command
    assert "echo $" not in lowered_build_command


def test_manual_migration_is_non_destructive():
    migration = (
        Path(app_module.BASE_DIR)
        / "migrations"
        / "001_auth_session_and_encrypted_api_keys.sql"
    ).read_text(encoding="utf-8")
    upper_migration = migration.upper()
    assert "CREATE TABLE IF NOT EXISTS" in upper_migration
    assert "DROP " not in upper_migration
    assert "TRUNCATE " not in upper_migration
    assert "DELETE FROM" not in upper_migration

    init_db_source = py_inspect.getsource(app_module.init_db).upper()
    assert "CHECKFIRST=TRUE" in init_db_source
    assert ".DROP(" not in init_db_source
    assert ".DROP_ALL(" not in init_db_source
    assert "TRUNCATE " not in init_db_source
    assert ".DELETE(" not in init_db_source


def test_production_import_does_not_run_database_schema_changes():
    child_env = os.environ.copy()
    child_env.update(
        {
            "FLASK_ENV": "production",
            "FLASK_SECRET_KEY": "production-import-test-secret",
            "API_KEY_ENCRYPTION_KEY": os.environ["API_KEY_ENCRYPTION_KEY"],
            "DATABASE_URL": "postgresql://user:pass@127.0.0.1:1/unreachable",
            "SESSION_TYPE": "sqlalchemy",
        }
    )
    child_env.pop("OPENAI_API_KEY", None)
    result = subprocess.run(
        [sys.executable, "-c", "import app; print('startup-ok')"],
        cwd=app_module.BASE_DIR,
        env=child_env,
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert result.stdout.strip() == "startup-ok"


@pytest.mark.filterwarnings(
    "ignore:datetime.datetime.utcnow\\(\\) is deprecated:DeprecationWarning"
)
def test_migrated_sqlalchemy_session_backend_waits_for_manual_schema(tmp_path):
    mini_app = Flask("session-test")
    mini_app.config.update(
        SECRET_KEY="session-test-secret",
        SQLALCHEMY_DATABASE_URI=f"sqlite:///{(tmp_path / 'session.db').as_posix()}",
        SESSION_PERMANENT=True,
    )
    mini_db = SQLAlchemy(mini_app)
    interface = app_module.MigratedSqlAlchemySessionInterface(mini_app, mini_db)
    mini_app.session_interface = interface

    with mini_app.app_context():
        assert "sessions" not in inspect(mini_db.engine).get_table_names()
        interface.sql_session_model.__table__.create(
            bind=mini_db.engine,
            checkfirst=True,
        )

    @mini_app.get("/set")
    def set_session_value():
        flask_session["marker"] = "persisten"
        return "ok"

    @mini_app.get("/get")
    def get_session_value():
        return flask_session.get("marker", "")

    mini_client = mini_app.test_client()
    assert mini_client.get("/set").status_code == 200
    assert mini_client.get("/get").data == b"persisten"


def test_local_secret_files_are_ignored_by_git():
    gitignore = (Path(app_module.BASE_DIR) / ".gitignore").read_text(
        encoding="utf-8"
    )
    assert ".env" in gitignore
    tracked_env = subprocess.run(
        ["git", "ls-files", "--error-unmatch", ".env"],
        cwd=app_module.BASE_DIR,
        capture_output=True,
        text=True,
        check=False,
    )
    assert tracked_env.returncode != 0


def test_secret_key_is_stable_and_required_in_production(monkeypatch):
    monkeypatch.setattr(app_module, "IS_PRODUCTION", True)
    monkeypatch.setenv("FLASK_SECRET_KEY", "secret-yang-tetap")
    monkeypatch.delenv("SECRET_KEY", raising=False)
    assert app_module.get_secret_key() == "secret-yang-tetap"
    assert app_module.get_secret_key() == "secret-yang-tetap"

    monkeypatch.delenv("FLASK_SECRET_KEY")
    with pytest.raises(RuntimeError, match="FLASK_SECRET_KEY"):
        app_module.get_secret_key()


def test_secret_key_stays_same_across_process_restarts(tmp_path):
    child_env = os.environ.copy()
    child_env.update(
        {
            "FLASK_ENV": "development",
            "FLASK_SECRET_KEY": "secret-restart-yang-sama",
            "SECRET_KEY": "",
            "SESSION_TYPE": "cachelib",
            "DATABASE_PATH": str(tmp_path / "restart.db"),
            "API_KEY_ENCRYPTION_KEY": os.environ["API_KEY_ENCRYPTION_KEY"],
        }
    )
    child_env.pop("OPENAI_API_KEY", None)
    child_env.pop("DATABASE_URL", None)
    child_env.pop("RENDER", None)
    command = [
        sys.executable,
        "-c",
        "import app; print(app.app.config['SECRET_KEY'])",
    ]
    first_start = subprocess.check_output(
        command,
        cwd=app_module.BASE_DIR,
        env=child_env,
        text=True,
    ).strip()
    second_start = subprocess.check_output(
        command,
        cwd=app_module.BASE_DIR,
        env=child_env,
        text=True,
    ).strip()
    assert first_start == second_start == "secret-restart-yang-sama"


# Health check tidak membocorkan informasi.
def test_healthz_public_and_minimal(client, monkeypatch):
    def unexpected_analysis():
        raise AssertionError("/healthz tidak boleh menjalankan analisis OpenAI")

    monkeypatch.setattr(app_module, "test_openai_connection", unexpected_analysis)
    response = client.get("/healthz")
    assert response.status_code == 200
    assert json.loads(response.data) == {"status": "ok"}


# Halaman sensitif memakai no-store.
def test_no_store_headers(client):
    register_and_login(client)
    response = client.get("/api-key")
    assert "no-store" in response.headers.get("Cache-Control", "")
